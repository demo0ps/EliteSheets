import hashlib
import json
from pathlib import Path

import pytest
import sqlalchemy as sa
from fastapi.testclient import TestClient
from cryptography.fernet import Fernet

from sheetqms.online import create_online_app
from sheetqms.cloud_store import events, knowledge
from sheetqms.security import password_hash, password_matches


@pytest.fixture
def client(tmp_path):
    app = create_online_app(tmp_path / 'preview', demo=True)
    with TestClient(app) as c:
        yield c, app


def sign(client, role='operator'):
    c, app = client
    response = c.post('/api/auth/demo', json={'role': role})
    assert response.status_code == 200
    value = response.json()
    c.headers.update({'X-Demo-Session': value['demo_token'], 'X-CSRF-Token': value['csrf']})
    if role == 'admin':
        c.headers['X-Workspace-ID'] = 'demo-a'
    else:
        c.headers.pop('X-Workspace-ID', None)
    return value


def example(client, case='pivot', request='example-request-001'):
    c, _ = client
    r = c.post('/api/demo/jobs', json={'case': case, 'request_id': request})
    assert r.status_code == 200, r.text
    return r.json()






def test_csrf_and_workspace_spoofing_are_blocked(client):
    c, _ = client
    sign(client)
    c.headers['X-CSRF-Token'] = 'wrong'
    assert c.post('/api/demo/jobs', json={'case': 'pivot', 'request_id': 'csrf-test-request'}).status_code == 403
    sign(client)
    c.headers['X-Workspace-ID'] = 'demo-b'
    assert c.get('/api/overview').status_code == 403




def test_released_workbook_and_evidence(client):
    c, _ = client
    sign(client)
    job = example(client)
    response = c.post(f'/api/jobs/{job["id"]}/run', json={})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['report']['status'] == 'RELEASED'
    assert result['report']['rows'][0]['Total'] == 50
    assert 'QMS released' in result['chat'][-1]['text']
    data = c.get(f'/api/jobs/{job["id"]}/output')
    assert data.status_code == 200 and data.content[:2] == b'PK'
    evidence = c.get(f'/api/jobs/{job["id"]}/evidence').json()
    assert evidence['audit_events'] and evidence['chat']
    assert 'object_key' not in json.dumps(result)
    assert 'path' not in result
















def test_preferences_rejects_secret_or_budget_fields(client):
    c, _ = client
    sign(client)
    assert c.put('/api/preferences', json={'theme': 'dark', 'api_key': 'FAKE_NOT_REAL'}).status_code == 422
    assert c.put('/api/preferences', json={'monthly_budget_usd': '100'}).status_code == 422
    assert c.put('/api/preferences', json={'theme': 'light', 'density': 'compact'}).status_code == 200
    assert c.get('/api/preferences').json()['theme'] == 'light'






def test_passwords_are_hashed_and_real_login_works(client):
    c, app = client
    sign(client, 'admin')
    c.post('/api/admin/customers', json={'name': 'Login test', 'email': 'real@example.test', 'display_name': 'Test', 'password': 'private-password-123'})
    value = app.state.store.user_email('real@example.test')
    assert value['password_hash'] != 'private-password-123'
    assert value['password_hash'].startswith('scrypt$')
    result = c.post('/api/auth/login', json={'email': 'real@example.test', 'password': 'private-password-123'})
    assert result.status_code == 200
    assert result.json()['user']['role'] == 'operator'
    assert c.post('/api/auth/login', json={'email': 'real@example.test', 'password': 'wrong'}).status_code == 401




def test_budget_reservations_are_atomic_integer_and_zero_blocks(client):
    _, app = client
    store = app.state.store
    with pytest.raises(ValueError):
        store.reserve('zero-budget', 1, 0)
    store.reserve('first-cost', 100, 200)
    with pytest.raises(ValueError):
        store.reserve('second-cost', 101, 200)
    assert store.spent() == 100
    store.settle('first-cost', 50)
    assert store.spent() == 50




def test_hash_gates_reject_modified_cloud_blob(client):
    c, app = client
    sign(client)
    job = example(client)
    c.post(f'/api/jobs/{job["id"]}/run', json={})
    value = app.state.store.job('demo-a', job['id'])
    app.state.blobs.put(value['output_object'], b'changed')
    assert c.get(f'/api/jobs/{job["id"]}/output').status_code == 409


def test_delete_removes_files_and_customer_job(client):
    c, app = client
    sign(client)
    job = example(client)
    stored = app.state.store.job('demo-a', job['id'])
    key = stored['files']['template']['object_key']
    assert c.delete(f'/api/jobs/{job["id"]}').status_code == 200
    assert c.get(f'/api/jobs/{job["id"]}').status_code == 404
    with pytest.raises(FileNotFoundError):
        app.state.blobs.get(key)


def test_production_rejects_ephemeral_sqlite_and_missing_cloud_config(tmp_path, monkeypatch):
    monkeypatch.setenv('SHEETQMS_ENV', 'production')
    with pytest.raises(ValueError, match='PostgreSQL'):
        create_online_app(tmp_path / 'invalid', demo=False, database_url='sqlite:///' + str(tmp_path / 'bad.db'))












def test_normal_workbook_without_rules_tab_and_header_dropdown_mapping(client):
    from io import BytesIO
    from openpyxl import Workbook
    c, _ = client
    sign(client)

    # Normal template WITHOUT a Rules sheet and with a multi-cell banner on Row 1
    tpl_wb = Workbook()
    tpl_ws = tpl_wb.active
    tpl_ws.title = 'Inventory_Report'
    tpl_ws.append(['ACME CORP — STANDARD INVENTORY TEMPLATE', None, 'Prepared: 2026-10-13'])
    tpl_ws.append([])
    tpl_ws.append(['SKU', 'Product Name', 'Qty', 'Unit Price'])
    tpl_bio = BytesIO()
    tpl_wb.save(tpl_bio)

    # Messy source with banner row, extra text column (Warehouse_Bin), two competing numeric columns, and Subtotal row
    src_wb = Workbook()
    src_ws = src_wb.active
    src_ws.title = 'Raw_ERP_Dump'
    src_ws.append(['ERP SYSTEM EXPORT — CONFIDENTIAL', None, None, 'Date: 2026-10-13'])
    src_ws.append([])
    src_ws.append(['Item code', 'Material Description', 'Warehouse_Bin', 'Crates_Delivered', 'Backup_Boxes', 'Unit Cost'])
    src_ws.append(['SKU-901', 'Stainless Flange 4in', 'Bin-A1', 24, 2, 45])
    src_ws.append(['SKU-902', 'Hydraulic Seal Kit', 'Bin-B4', 15, 0, 18])
    src_ws.append(['Subtotal', 'All Bins', '', 39, 2, 63])
    src_bio = BytesIO()
    src_wb.save(src_bio)

    job = c.post('/api/jobs', data={'request_id': 'no-rules-tab-001'}, files=[
        ('target', ('Normal-Template.xlsx', tpl_bio.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')),
        ('sources', ('Messy-ERP.xlsx', src_bio.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')),
    ]).json()

    # Banner rows skipped, Subtotal row excluded (2 rows), and text column Warehouse_Bin never misplaced into Qty
    assert job['plan']['sources'][0]['field_map'] == {'SKU': 'Item code', 'Product Name': 'Material Description', 'Unit Price': 'Unit Cost'}
    assert len(job['report']['rows']) == 2
    assert job['state'] == 'HELD'

    # Selecting Crates_Delivered from the Qty column header dropdown (/map-column) immediately maps and releases
    mapped = c.post(f'/api/jobs/{job["id"]}/map-column', json={'target_column': 'Qty', 'source_column': 'Crates_Delivered', 'request_id': 'map-col-001'}).json()
    assert mapped['state'] == 'RELEASED'
    assert mapped['report']['rows'][0] == {'SKU': 'SKU-901', 'Product Name': 'Stainless Flange 4in', 'Qty': 24, 'Unit Price': 45}


def test_sasams_two_sheet_student_and_subject_percentage_matrix_import(client):
    from io import BytesIO
    from openpyxl import load_workbook
    c, _ = client
    sign(client)
    job = example(client, 'sasams_students', 'sasams-students-001')
    ran = c.post(f'/api/jobs/{job["id"]}/run', json={}).json()
    assert ran['state'] == 'RELEASED'
    assert ran['report']['headers'] == [
        'Student ID', 'First Name', 'Last Name', 'Date of Birth', 'Grade',
        'Subjects', 'Parent First Name', 'Parent Last Name', 'Parent Email', 'Parent Phone',
    ]
    rows = ran['report']['rows']
    assert len(rows) == 4
    assert rows[0] == {
        'Student ID': 'STU001', 'First Name': 'John', 'Last Name': 'Doe',
        'Date of Birth': '2015-03-15', 'Grade': 'Grade 8',
        'Subjects': 'Mathematics; English; Science',
        'Parent First Name': 'Jane', 'Parent Last Name': 'Doe',
        'Parent Email': 'jane.doe@example.com', 'Parent Phone': '0821234567',
    }
    assert rows[1] == {
        'Student ID': 'STU002', 'First Name': 'Sarah', 'Last Name': 'Smith',
        'Date of Birth': '2014-07-22', 'Grade': 'Grade 9',
        'Subjects': 'Mathematics; English; Physical Science',
        'Parent First Name': 'Michael', 'Parent Last Name': 'Smith',
        'Parent Email': 'michael.smith@example.com', 'Parent Phone': '0839876543',
    }
    assert rows[2] == {
        'Student ID': 'STU003', 'First Name': 'Emma', 'Last Name': 'Johnson',
        'Date of Birth': '2016-01-10', 'Grade': 'Grade 7',
        'Subjects': 'Mathematics; English; Natural Science',
        'Parent First Name': 'Patricia', 'Parent Last Name': 'Johnson',
        'Parent Email': 'patricia.johnson@example.com', 'Parent Phone': '0847654321',
    }
    assert rows[3] == {
        'Student ID': 'STU004', 'First Name': 'Liam', 'Last Name': 'Johnson',
        'Date of Birth': '2013-09-05', 'Grade': 'Grade 10',
        'Subjects': 'Mathematics; English; Physical Science; Geography',
        'Parent First Name': 'Patricia', 'Parent Last Name': 'Johnson',
        'Parent Email': 'patricia.johnson@example.com', 'Parent Phone': '0847654321',
    }

    # Verify exported .xlsx preserves leading zero on Parent Phone ('0821234567') and exact cell values
    dl = c.get(f'/api/jobs/{job["id"]}/output')
    assert dl.status_code == 200
    wb = load_workbook(BytesIO(dl.content), data_only=True)
    ws = wb['Students']
    assert ws['A2'].value == 'STU001'
    assert ws['E2'].value == 'Grade 8'
    assert ws['F2'].value == 'Mathematics; English; Science'
    assert ws['J2'].value == '0821234567'
    assert ws['F5'].value == 'Mathematics; English; Physical Science; Geography'


















def test_grid_edits_are_saved_as_values_and_rebuild_the_workbook(client):
    # Elite Sheets batch 1: Excel-style grid saves values only; a typed '=' stays plain text.
    from io import BytesIO
    from openpyxl import load_workbook
    c, _ = client
    sign(client)
    job = example(client, request='grid-edit-001')
    ran = c.post(f'/api/jobs/{job["id"]}/run', json={})
    assert ran.status_code == 200, ran.text
    report = ran.json()['report']
    headers = report['headers']
    rows = [dict(r) for r in report['rows']]
    # Computed columns (sum rules) are recalculated on save, so the edit goes to a plain text column.
    rows[0][headers[0]] = 'Grid edit'
    rows[1][headers[0]] = '=SUM(A1:A2)'
    saved = c.put(f'/api/jobs/{job["id"]}/rows', json={'headers': headers, 'rows': rows})
    assert saved.status_code == 200, saved.text
    assert saved.json()['report']['rows'][0][headers[0]] == 'Grid edit'
    assert saved.json()['report']['rows'][1][headers[0]] == '=SUM(A1:A2)'
    out = c.get(f'/api/jobs/{job["id"]}/output')
    assert out.status_code == 200 and out.content[:2] == b'PK'
    ws = load_workbook(BytesIO(out.content)).active
    formula_cells = [cell for row in ws.iter_rows() for cell in row if cell.value == '=SUM(A1:A2)']
    assert formula_cells and all(cell.data_type == 's' for cell in formula_cells)
    assert 'Grid edit' in [cell.value for row in ws.iter_rows() for cell in row]
    # Rows that are not a list of objects are refused.
    bad = c.put(f'/api/jobs/{job["id"]}/rows', json={'headers': headers, 'rows': 'nope'})
    assert bad.status_code == 422


def test_anonymous_cannot_read_jobs_or_admin(client):
    c, _ = client
    assert c.get('/api/overview').status_code == 401
    assert c.get('/api/admin/settings').status_code == 401
    assert c.get('/').status_code == 200


def test_operator_cannot_read_or_change_platform_secrets(client):
    c, _ = client
    sign(client)
    assert c.get('/api/admin/settings').status_code == 403
    assert c.put('/api/admin/settings', json={'monthly_budget_usd': '100'}).status_code == 403
    assert c.post('/api/admin/customers', json={'name': 'Xx', 'email': 'x@example.test', 'display_name': 'Xx', 'password': 'password-12345'}).status_code == 403
    assert 'base_url' not in c.get('/api/overview').json()['ai']


def test_admin_can_create_amounts_and_customer_accounts(client):
    c, _ = client
    sign(client, 'admin')
    original = c.get('/api/admin/settings').json()
    assert original['platform_budget_usd'] == original['monthly_budget_usd'] == '0'
    saved = c.put('/api/admin/settings', json={'platform_budget_usd': '20', 'monthly_budget_usd': '5'}).json()
    assert saved['platform_budget_usd'] == '20'
    assert 'api_key' not in saved
    created = c.post('/api/admin/customers', json={'name': 'Fictional customer', 'email': 'test@example.test', 'display_name': 'Test operator', 'password': 'private-test-password'}).json()
    assert created['user']['role'] == 'operator'
    assert 'password' not in json.dumps(created)
    assert len(c.get('/api/admin/customers').json()['items']) == 3
