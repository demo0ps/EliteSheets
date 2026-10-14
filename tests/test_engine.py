import hashlib
from pathlib import Path
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from sheetqms.engine import infer_plan, transform, decimal_value
from sheetqms.workbooks import check_file
from conftest import make_job


def execute(samples, tmp_path, case='pivot'):
    job = make_job(samples, case)
    plan = infer_plan(job)
    output = tmp_path / 'output.xlsx'
    return job, plan, transform(job, plan, output), output


def test_row_to_column_known_answer(samples, tmp_path):
    job, plan, report, output = execute(samples, tmp_path)
    assert report.status == 'RELEASED', report.model_dump()
    assert report.rows == [{'Site': 'North', 'Jan': 12, 'Feb': 18, 'Mar': 20, 'Total': 50}, {'Site': 'South', 'Jan': 7, 'Feb': 15, 'Mar': 13, 'Total': 35}, {'Site': 'West', 'Jan': 9, 'Feb': 11, 'Mar': 15, 'Total': 35}]
    assert report.source_records == report.mapped_records == 9
    assert {s.layout for s in plan.sources} == {'long', 'matrix'}
    assert all(c.passed for c in report.checks)
    assert output.is_file()


def test_formula_style_and_cached_results_preserved(samples, tmp_path):
    _, _, report, output = execute(samples, tmp_path)
    assert report.status == 'RELEASED'
    wb = load_workbook(output, data_only=False)
    assert wb['Output']['E4'].value == '=SUM(B4:D4)'
    assert wb['Output']['E5'].value == '=SUM(B5:D5)'
    assert wb['Output']['A3'].fill.fgColor.rgb.endswith('3559E0')
    wb.close()
    wb = load_workbook(output, data_only=True)
    assert [wb['Output'][f'E{r}'].value for r in [4, 5, 6]] == [50, 35, 35]
    wb.close()


def test_wide_to_long_known_answer(samples, tmp_path):
    _, plan, report, _ = execute(samples, tmp_path, 'unpivot')
    assert report.status == 'RELEASED', report.model_dump()
    assert plan.sources[0].layout == 'wide'
    assert report.output_rows == report.source_records == 6
    assert report.rows[0] == {'Site': 'North', 'Period': 'Jan', 'Units': 12}
    assert report.rows[-1] == {'Site': 'South', 'Period': 'Mar', 'Units': 13}


def test_csv_direct_fields_are_typed_numbers(samples, tmp_path):
    _, plan, report, output = execute(samples, tmp_path, 'direct')
    assert report.status == 'RELEASED', report.model_dump()
    assert plan.sources[0].field_map == {'Item': 'Item code', 'Description': 'Name', 'Qty': 'Quantity'}
    assert report.rows[0]['Qty'] == 8
    wb = load_workbook(output, data_only=True)
    assert wb['Output']['C4'].value == 8
    wb.close()


def test_missing_is_held_not_invented_zero(samples, tmp_path):
    _, _, report, output = execute(samples, tmp_path, 'hold')
    assert report.status == 'HELD'
    assert report.rows[0]['Feb'] is None
    assert any('North / Feb' in q.text for q in report.questions)
    assert not output.exists()


def test_original_files_unchanged(samples, tmp_path):
    hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in samples.iterdir()}
    execute(samples, tmp_path)
    assert hashes == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in samples.iterdir()}


def test_duplicate_values_need_aggregation_rule(samples, tmp_path):
    path = samples / 'north-messy.xlsx'
    wb = load_workbook(path)
    wb.active.append(['North', 'Jan', 99])
    wb.save(path)
    _, _, report, output = execute(samples, tmp_path)
    assert report.status == 'HELD'
    assert any('Multiple values' in q.text for q in report.questions)
    assert not output.exists()


def test_negative_fails_native_and_business_checks(samples, tmp_path):
    path = samples / 'north-messy.xlsx'
    wb = load_workbook(path)
    wb.active['C5'] = -1
    wb.save(path)
    _, _, report, _ = execute(samples, tmp_path)
    assert report.status == 'HELD'
    assert any(not c.passed and 'nonnegative' in c.name for c in report.checks)


def test_fraction_fails_native_whole_number_validation(samples, tmp_path):
    path = samples / 'north-messy.xlsx'
    wb = load_workbook(path)
    wb.active['C5'] = 1.5
    wb.save(path)
    _, _, report, _ = execute(samples, tmp_path)
    assert report.status == 'HELD'
    assert any(not c.passed and 'integer' in c.name for c in report.checks)


def test_uncached_source_formula_held(samples, tmp_path):
    path = samples / 'north-messy.xlsx'
    wb = load_workbook(path)
    wb.active['C5'] = '=6+6'
    wb.save(path)
    _, _, report, output = execute(samples, tmp_path)
    assert report.status == 'HELD'
    assert any('cached result' in q.text for q in report.questions)
    assert not output.exists()


def test_unsupported_target_formula_held(samples, tmp_path):
    path = samples / 'monthly-output-template.xlsx'
    wb = load_workbook(path)
    wb['Output']['E4'] = '=ROUND(SUM(B4:D4),2)'
    wb.save(path)
    _, _, report, output = execute(samples, tmp_path)
    assert report.status == 'HELD'
    assert any(q.id == 'export-validation' for q in report.questions)
    assert not output.exists()


def test_no_silent_number_locale_guess():
    with pytest.raises(ValueError):
        decimal_value('1,234')
    assert decimal_value('1,234', {'thousands_separator': ','}) == Decimal(1234)
    assert decimal_value('1.234,50', {'thousands_separator': '.', 'decimal_separator': ','}) == Decimal('1234.50')


def test_csv_formula_injection_is_written_as_text(samples, tmp_path):
    path = samples / 'items-source.csv'
    path.write_text('Item code,Name,Quantity\nSKU-001,=1+1,8\n', encoding='utf-8')
    _, _, report, output = execute(samples, tmp_path, 'direct')
    assert report.status == 'RELEASED'
    wb = load_workbook(output, data_only=False)
    assert wb['Output']['B4'].value == '=1+1'
    assert wb['Output']['B4'].data_type == 's'
    wb.close()


def test_unknown_rule_cannot_be_silently_skipped(samples, tmp_path):
    path = samples / 'monthly-output-template.xlsx'
    wb = load_workbook(path)
    wb['Rules'].append(['R99', 'magic_tax_adjustment', 'Total', 'unknown'])
    wb.save(path)
    _, _, report, output = execute(samples, tmp_path)
    assert report.status == 'HELD'
    assert any(c.id == 'rule-coverage' and not c.passed for c in report.checks)
    assert not output.exists()


def test_ai_cannot_drop_authoritative_rules(samples, tmp_path):
    job = make_job(samples)
    plan = infer_plan(job)
    plan.rules = []
    report = transform(job, plan, tmp_path / 'output.xlsx')
    assert report.status == 'HELD'
    assert any(c.id == 'rule-integrity' and not c.passed for c in report.checks)


def test_source_regions_cannot_be_mapped_twice(samples, tmp_path):
    job = make_job(samples)
    plan = infer_plan(job)
    plan.sources.append(plan.sources[0].model_copy())
    report = transform(job, plan, tmp_path / 'output.xlsx')
    assert report.status == 'HELD'
    assert any('mapped twice' in q.text for q in report.questions)


def test_unknown_source_reference_is_held(samples, tmp_path):
    job = make_job(samples)
    plan = infer_plan(job)
    plan.sources[0].source_id = 'invented'
    report = transform(job, plan, tmp_path / 'output.xlsx')
    assert report.status == 'HELD'


def test_unmapped_extra_column_is_held(samples, tmp_path):
    path = samples / 'north-messy.xlsx'
    wb = load_workbook(path)
    wb.active['D4'] = 'Extra metric'
    for r in [5, 6, 7]:
        wb.active[f'D{r}'] = 123
    wb.save(path)
    _, _, report, _ = execute(samples, tmp_path)
    assert report.status == 'HELD'
    assert any('unmapped' in q.text.lower() for q in report.questions)


def test_unsupported_file_rejected(tmp_path):
    file = tmp_path / 'legacy.xls'
    file.write_text('not Excel')
    with pytest.raises(ValueError, match='supports'):
        check_file(file)


def test_circular_formula_held(samples, tmp_path):
    path = samples / 'monthly-output-template.xlsx'
    wb = load_workbook(path)
    wb['Output']['E4'] = '=E4'
    wb.save(path)
    _, _, report, _ = execute(samples, tmp_path)
    assert report.status == 'HELD'
    assert any('Circular' in q.text for q in report.questions)
