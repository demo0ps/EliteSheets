from decimal import Decimal

import pytest
from openpyxl import load_workbook

from sheetqms.gradebook import Assessment, Learner, build_rows, calculate, export_xlsx, validate


def term_setup():
    return [
        Assessment('Test 1', 6, 10),
        Assessment('Test 2', 80, 30),
        Assessment('Exam', 40, 60),
    ]


def errors(checks):
    return [c.message for c in checks if c.level == 'error']


def warnings(checks):
    return [c.message for c in checks if c.level == 'warning']


def test_percentages_are_taken_before_weighting():
    # 6/6 on a small test should count the same as a full mark on a large test.
    a = [Assessment('Small', 6, 50), Assessment('Big', 80, 50)]
    result = calculate(a, Learner('A1', {'Small': 6, 'Big': 0}))
    assert result['percentages']['Small'] == Decimal('100.00')
    assert result['term_mark'] == Decimal('50.00')


def test_weighted_term_mark_and_caps_level():
    a = term_setup()
    learner = Learner('A1', {'Test 1': 3, 'Test 2': 60, 'Exam': 20})
    result = calculate(a, learner)
    # 50% x 10 + 75% x 30 + 50% x 60 = 5 + 22.5 + 30 = 57.5
    assert result['term_mark'] == Decimal('57.50')
    assert result['caps_level'] == 4  # 50-59.99 is level 4


@pytest.mark.parametrize('mark,level', [(29.99, 1), (30, 2), (49.99, 3), (50, 4), (69.99, 5), (70, 6), (80, 7)])
def test_caps_level_boundaries(mark, level):
    a = [Assessment('Only', 100, 100)]
    assert calculate(a, Learner('A1', {'Only': mark}))['caps_level'] == level


def test_no_term_mark_when_a_weighted_mark_is_missing():
    a = term_setup()
    result = calculate(a, Learner('A1', {'Test 1': 3, 'Exam': 20}))
    assert result['term_mark'] is None
    assert result['caps_level'] is None


def test_absent_counts_as_missing_for_term_mark():
    a = term_setup()
    result = calculate(a, Learner('A1', {'Test 1': 3, 'Test 2': 'absent', 'Exam': 20}))
    assert result['term_mark'] is None


def test_weights_must_total_100():
    a = [Assessment('A', 10, 40), Assessment('B', 10, 40)]
    assert any('80%' in m for m in errors(validate(a, [])))


def test_mark_above_maximum_is_an_error():
    a = term_setup()
    checks = validate(a, [Learner('A1', {'Test 1': 7, 'Test 2': 60, 'Exam': 20})])
    assert any('above the maximum' in m for m in errors(checks))


def test_negative_and_non_numeric_marks_are_errors():
    a = term_setup()
    checks = validate(a, [Learner('A1', {'Test 1': -1, 'Test 2': 'abc', 'Exam': 20})])
    msgs = errors(checks)
    assert any('negative' in m for m in msgs)
    assert any('not a number' in m for m in msgs)


def test_duplicate_learner_numbers_are_errors():
    a = term_setup()
    learners = [Learner('A1', {}), Learner('A1', {})]
    assert any('more than once' in m for m in errors(validate(a, learners)))


def test_zero_maximum_is_an_error():
    a = [Assessment('Bad', 0, 100)]
    assert any('maximum mark greater than zero' in m for m in errors(validate(a, [])))


def test_missing_mark_is_a_warning_not_an_error():
    a = term_setup()
    checks = validate(a, [Learner('A1', {'Test 1': 3, 'Exam': 20})])
    assert not errors(checks)
    assert any('no mark for "Test 2"' in m for m in warnings(checks))


def test_build_rows_shows_blank_when_incomplete():
    a = term_setup()
    rows = build_rows(a, [Learner('A1', {'Test 1': 3}), Learner('A2', {'Test 1': 6, 'Test 2': 80, 'Exam': 40})])
    assert rows[0]['Term mark %'] == ''
    assert rows[1]['Term mark %'] == 100.0
    assert rows[1]['CAPS level'] == 7


def test_export_blocked_when_errors(tmp_path):
    a = term_setup()
    with pytest.raises(ValueError):
        export_xlsx(tmp_path / 'bad.xlsx', a, [Learner('A1', {'Test 1': 99, 'Test 2': 1, 'Exam': 1})])
    assert not (tmp_path / 'bad.xlsx').exists()


def test_export_writes_marks_and_checks(tmp_path):
    a = term_setup()
    learners = [Learner('A1', {'Test 1': 3, 'Test 2': 60, 'Exam': 20}), Learner('A2', {'Test 1': 6})]
    path = tmp_path / 'class.xlsx'
    checks = export_xlsx(path, a, learners)
    wb = load_workbook(path)
    assert wb.sheetnames == ['Marks', 'Checks']
    marks = list(wb['Marks'].values)
    assert marks[0][0] == 'Learner'
    assert marks[1][0] == 'A1'
    assert marks[1][-2] == 57.5
    assert any(c.level == 'warning' for c in checks)
