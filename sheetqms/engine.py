from __future__ import annotations

import ast
import copy
import hashlib
import re
import zipfile
from collections import OrderedDict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries

from .models import MappingPlan, DestinationPlan, SourcePlan, Rule, Question, Check, Report
from .store import canonical, digest
from .workbooks import book_rows, norm, public_value, builtin_aliases_for, PERIOD_TOKENS


MAX_NORMALIZED_RECORDS = 50000
MAX_OUTPUT_ROWS = 20000
MAX_OUTPUT_CELLS = 1_000_000


def col_sample_type(preview_rows, col_idx: int) -> str:
    vals = [r[col_idx] for r in (preview_rows or []) if col_idx < len(r) and r[col_idx] is not None and str(r[col_idx]).strip() != '']
    if not vals:
        return 'any'
    num_count = 0
    code_count = 0
    period_count = 0
    for v in vals:
        if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
            num_count += 1
            continue
        s = str(v).strip()
        if re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)', s):
            num_count += 1
        elif norm(s) in PERIOD_TOKENS:
            period_count += 1
        elif re.fullmatch(r'[A-Za-z0-9]+[-_][A-Za-z0-9-_]+', s) or (len(s) <= 12 and any(c.isdigit() for c in s) and ' ' not in s):
            code_count += 1
    if num_count >= len(vals) * 0.7:
        return 'numeric'
    if period_count >= len(vals) * 0.7:
        return 'period'
    if code_count >= len(vals) * 0.7:
        return 'code'
    return 'text'


def expected_field_type(field: str, rules) -> str:
    nf = norm(field)
    tokens = set(re.findall(r'[a-z0-9]+', nf))
    for r in rules:
        rk = r['kind'] if isinstance(r, dict) else r.kind
        rfields = r['fields'] if isinstance(r, dict) else r.fields
        if rk in {'numeric', 'integer', 'nonnegative', 'min', 'max', 'sum'} and field in rfields:
            return 'numeric'
    if nf in PERIOD_TOKENS or tokens & {'qty', 'quantity', 'units', 'count', 'price', 'cost', 'rate', 'total', 'amount', 'revenue', 'sales', 'value', 'sum', 'boxes', 'crates', 'weight', 'score'}:
        return 'numeric'
    if tokens & {'item', 'sku', 'code', 'id', 'part', 'ref'}:
        return 'code'
    if tokens & {'description', 'name', 'title', 'details', 'notes', 'site', 'branch', 'location', 'store', 'customer', 'client', 'supplier', 'vendor', 'category', 'status', 'type', 'department'}:
        return 'text'
    return 'any'


def compatible_type(expected: str, actual: str) -> bool:
    if expected == 'any' or actual == 'any':
        return True
    if expected == 'numeric':
        return actual == 'numeric'
    if expected in {'text', 'code'}:
        return actual != 'numeric'
    return True


def question(qid, text, evidence='', kind='mapping', options=None):
    return Question(id=qid, text=text, evidence=evidence, kind=kind, options=options or [])


def rules_for_job(job):
    result = list(job['profiles']['template']['rules'])
    for source in job['profiles']['sources']:
        for rule in source['rules']:
            result.append(rule | {'id': source['id'] + ':' + rule['id'], 'evidence': source['id'] + ':' + rule['evidence']})
    return result


def unknown_for_job(job):
    items = list(job['profiles']['template']['unknown_rules'])
    for source in job['profiles']['sources']:
        items += [item | {'ref': source['id'] + ':' + item['ref']} for item in source['unknown_rules']]
    return items


def aliases_for(template, field):
    base = [field] + template['aliases'].get(field, []) + builtin_aliases_for(field)
    return {norm(x) for x in base if x}


def find(headers, aliases):
    return [header for header in headers if norm(header) in aliases]


def finalise_cell_profiles(job, plan):
    for source in plan.sources:
        if source.layout != 'cells':
            continue
        file = next(f for f in job['files']['sources'] if f['id'] == source.source_id)
        rows = book_rows(Path(file['path']))[source.sheet]
        first_label = min((int(re.search(r'\d+', cell).group()) for cell in source.field_labels.values() if re.fullmatch(r'A[1-9]\d*', cell)), default=1)
        labels = [(r, str(row[0])) for r, row in enumerate(rows, 1) if r >= first_label and row and row[0] is not None]
        profile = next(p for p in job['profiles']['sources'] if p['id'] == source.source_id)
        sheet = next(s for s in profile['sheets'] if s['name'] == source.sheet)
        sheet['layout_fingerprint'] = digest({'sheet': source.sheet, 'kind': 'row-label record', 'labels': labels})


def normalise_cell_form(source, job, raw_rows, rows, numeric_fields, config):
    problems, exclusions, refs, data = [], [], [], {}
    used = set()
    profile = next(p for p in job['profiles']['sources'] if p['id'] == source.source_id)
    sheet = next(s for s in profile['sheets'] if s['name'] == source.sheet)
    formula_cells = set(sheet['formula_cells'])
    def read(coordinate):
        if not re.fullmatch(r'[AB][1-9]\d*', coordinate):
            raise ValueError('Row-label mode supports labels in A and values in B; other forms require an added, verified operation.')
        col, number, _, _ = range_boundaries(coordinate + ':' + coordinate)
        if number > len(rows) or col > len(rows[number - 1]):
            raise ValueError('Cell reference is outside the source region.')
        if coordinate in formula_cells:
            raise ValueError('Formula-valued row-label cells need a verified cell-form formula operation; this version holds them.')
        return rows[number - 1][col - 1]
    for field, coordinate in source.field_map.items():
        label_ref = source.field_labels.get(field, '')
        try:
            if not re.fullmatch(r'B[1-9]\d*', coordinate) or label_ref != 'A' + coordinate[1:]:
                raise ValueError('Each row-label mapping needs an adjacent A-label / B-value pair in the same row.')
            label, value = read(label_ref), read(coordinate)
            if norm(label) not in aliases_for(job['profiles']['template'], field) and not job.get('plan_confirmation'):
                raise ValueError(f'Label {label!s} is not an authorised alias for {field}; ask for confirmation.')
            data[field] = decimal_value(value, config) if field in numeric_fields else numeric_or_original(value)
            used.update([label_ref, coordinate])
            refs.append({'source_id': source.source_id, 'sheet': source.sheet, 'cells': [label_ref, coordinate], 'target_field': field})
        except (ValueError, InvalidOperation) as exc:
            problems.append(question('cell-field-' + digest([source.source_id, source.sheet, field])[:10], str(exc), source.sheet + '!' + coordinate, kind='data'))
    first_mapped_row = min((int(c[1:]) for c in used), default=1)
    for number, row in enumerate(raw_rows, 1):
        filled = [(i, value) for i, value in enumerate(row, 1) if value is not None and str(value).strip()]
        header_pair = len(row) >= 2 and norm(row[0]) in {'field', 'label', 'attribute', 'key'} and norm(row[1]) in {'value', 'data', 'content'}
        title = number < first_mapped_row and len(filled) == 1 and isinstance(filled[0][1], str)
        for col, value in filled:
            coordinate = f'{get_column_letter(col)}{number}'
            if coordinate in used:
                continue
            if header_pair or title:
                exclusions.append({'source_id': source.source_id, 'sheet': source.sheet, 'cell': coordinate, 'reason': 'Recognised form header or single-cell pre-label title.'})
            elif coordinate in source.ignored_cells:
                reason = source.ignored_cells[coordinate].strip()
                if not reason:
                    problems.append(question('cell-exclusion-reason', 'Excluded cells need a reason.', coordinate))
                exclusions.append({'source_id': source.source_id, 'sheet': source.sheet, 'cell': coordinate, 'reason': reason})
            else:
                problems.append(question('cell-scope-' + digest([source.source_id, source.sheet, coordinate])[:10], 'This source cell is outside the mapping. Clarify its destination or explicitly approve exclusion.', source.sheet + '!' + coordinate, kind='mapping'))
    if source.ignored_columns:
        problems.append(question('form-columns', 'Use explicit ignored_cells for a row-label form, not column-wide exclusions.', source.sheet))
    record = {'id': f'{source.source_id}:{source.sheet}:form', 'data': data, 'refs': refs, 'operation': 'row labels into column fields'}
    return record, problems, exclusions


def input_signature(job):
    return digest(sorted([digest([s['layout_fingerprint'] for s in p['sheets']]) for p in job['profiles']['sources']]))


def infer_plan(job):
    template = job['profiles']['template']
    config = template['config']
    sheet = next(s for s in template['sheets'] if s['name'] == template['output_sheet'])
    headers = sheet['headers']
    is_auto_template = any(r.get('id', '').startswith('AUTO-') for r in template['rules'])
    problems = []
    mode = config.get('mode', '')
    if mode not in {'pivot', 'long', 'direct'}:
        problems.append(question('destination-mode', 'What output shape is required: pivoted columns, long rows, or direct field mapping?', sheet['name'], options=['pivot', 'long', 'direct']))
        mode = 'direct'
    key = config.get('key') or (headers[0] if headers else '')
    if not config.get('key'):
        problems.append(question('destination-key', 'Which destination field identifies each record?', sheet['name']))
    categories = [x.strip() for x in config.get('categories', '').split('|') if x.strip()]
    if mode in {'pivot', 'long'} and not categories:
        problems.append(question('categories', 'Which period/category labels are allowed? Add categories to the embedded rules.', sheet['name'], kind='rule'))
    destination = DestinationPlan(sheet=sheet['name'], header_row=sheet['header_row'], mode=mode, key=key, period=config.get('period', ''), value=config.get('value', ''), categories=categories, total=config.get('total', ''))
    sources = []
    key_aliases = aliases_for(template, key)
    period_aliases = aliases_for(template, destination.period)
    value_aliases = aliases_for(template, destination.value)
    direct_candidates = []
    for profile in job['profiles']['sources']:
        for source_sheet in profile['sheets']:
            candidates = source_sheet['candidates'] or [{'row': source_sheet['header_row'], 'headers': source_sheet['headers'], 'score': 0}]
            scored = []
            for candidate in candidates:
                hs = candidate['headers']
                if mode == 'direct':
                    hits = sum(bool(find(hs, aliases_for(template, field))) for field in headers)
                else:
                    hits = len(find(hs, key_aliases)) + len(find(hs, period_aliases)) + len(find(hs, value_aliases)) + sum(norm(c) in {norm(h) for h in hs} for c in categories)
                scored.append((hits, candidate['score'], candidate))
            best = max(scored, key=lambda item: (item[0], item[1]))[2]
            hs = list(best['headers'])
            while hs and not hs[-1]:
                hs.pop()
            preview = source_sheet.get('preview', [])
            col_types = {h: col_sample_type(preview, i) for i, h in enumerate(hs)}
            prefix = f'{profile["id"]}:{source_sheet["name"]}'
            if not hs or any(not h for h in hs) or len({norm(h) for h in hs}) != len(hs):
                problems.append(question(prefix + ':headers', 'This region has blank or duplicated headings. Confirm the header row and uniquely identify its fields.', f'{profile["name"]} / {source_sheet["name"]}'))
                continue
            if mode == 'direct':
                field_map = {}
                form_rows = book_rows(Path(next(f['path'] for f in job['files']['sources'] if f['id'] == profile['id'])))[source_sheet['name']]
                cell_fields, cell_labels = {}, {}
                for field in headers:
                    aliases = aliases_for(template, field)
                    matches = [(r, row[0]) for r, row in enumerate(form_rows, 1) if len(row) >= 2 and norm(row[0]) in aliases]
                    if len(matches) == 1:
                        r = matches[0][0]
                        cell_fields[field], cell_labels[field] = f'B{r}', f'A{r}'
                if len(cell_fields) == len(headers):
                    sources.append(SourcePlan(source_id=profile['id'], sheet=source_sheet['name'], header_row=best['row'], layout='cells', field_map=cell_fields, field_labels=cell_labels))
                    continue
                match_counts = {}
                for field in headers:
                    exp_t = expected_field_type(field, template['rules'])
                    raw_matches = [m for m in find(hs, aliases_for(template, field)) if m not in field_map.values()]
                    matches = [m for m in raw_matches if compatible_type(exp_t, col_types.get(m, 'any'))]
                    if not matches and raw_matches and any(norm(m) == norm(field) for m in raw_matches):
                        matches = [m for m in raw_matches if norm(m) == norm(field)]
                    if not matches and is_auto_template:
                        ftoks = set(re.findall(r'[a-z0-9]+', norm(field))) - {'the', 'and', 'for', 'of', 'in', 'to'}
                        if ftoks:
                            matches = [h for h in hs if h not in field_map.values() and compatible_type(exp_t, col_types.get(h, 'any')) and bool(ftoks & set(re.findall(r'[a-z0-9]+', norm(h))))]
                    match_counts[field] = len(matches)
                    if len(matches) == 1:
                        field_map[field] = matches[0]
                # Detect SA-SAMS / horizontal Subject-Percentage matrix columns for a list target field (e.g. 'Subjects')
                list_field, list_cols = '', []
                list_targets = [f for f in headers if norm(f) in {'subjects', 'subject', 'enrolled subjects', 'subject list', 'courses', 'modules'} and f not in field_map]
                if list_targets:
                    non_list_aliases = {a for f in headers if f not in list_targets for a in aliases_for(template, f)}
                    skip_meta = {'total', 'subtotal', 'grand total', 'average', 'grade average', 'class', 'gender', 'term', 'age', 'rank', 'pass', 'fail', 'result', 'comments', 'notes'}
                    cand_subj = [h for h in hs if h not in field_map.values() and norm(h) not in non_list_aliases and norm(h) not in skip_meta]
                    if len(cand_subj) >= 2:
                        list_field = list_targets[0]
                        list_cols = cand_subj
                direct_candidates.append({
                    'prefix': prefix, 'profile_id': profile['id'], 'profile_name': profile['name'],
                    'sheet_name': source_sheet['name'], 'header_row': best['row'], 'hs': hs,
                    'field_map': field_map, 'match_counts': match_counts,
                    'list_field': list_field, 'list_cols': list_cols
                })
            else:
                keys, periods, values = find(hs, key_aliases), find(hs, period_aliases), find(hs, value_aliases)
                category_headers = [h for h in hs if norm(h) in {norm(c) for c in categories}]
                if not keys and len(category_headers) == len(categories):
                    cand_keys = [h for h in hs if h not in category_headers and norm(h) not in {'total', 'subtotal', 'grand total', 'sum'} and col_types.get(h) != 'numeric']
                    if len(cand_keys) == 1:
                        keys = cand_keys
                if not periods and not keys and len(hs) > 1 and col_types.get(hs[0]) == 'period':
                    periods = [hs[0]]
                if len(keys) == len(periods) == len(values) == 1:
                    source = SourcePlan(source_id=profile['id'], sheet=source_sheet['name'], header_row=best['row'], layout='long', key=keys[0], period=periods[0], value=values[0])
                    used = [keys[0], periods[0], values[0]]
                elif len(keys) == 1 and len(category_headers) == len(categories):
                    ignored_cols = {}
                    for h in hs:
                        if h not in ([keys[0]] + category_headers) and (is_auto_template or norm(h) in {'total', 'subtotal', 'grand total'}):
                            ignored_cols[h] = 'Auto-excluded summary/extra column'
                    source = SourcePlan(source_id=profile['id'], sheet=source_sheet['name'], header_row=best['row'], layout='wide', key=keys[0], columns=category_headers, ignored_columns=ignored_cols)
                    used = [keys[0]] + category_headers + list(ignored_cols.keys())
                elif len(periods) == 1 and not keys and len(hs) > 1:
                    entity_columns = [h for h in hs if h != periods[0] and norm(h) not in {'total', 'subtotal', 'grand total'}]
                    ignored_cols = {h: 'Auto-excluded summary column' for h in hs if h != periods[0] and h not in entity_columns}
                    source = SourcePlan(source_id=profile['id'], sheet=source_sheet['name'], header_row=best['row'], layout='matrix', period=periods[0], columns=entity_columns, ignored_columns=ignored_cols)
                    used = hs
                else:
                    problems.append(question(prefix + ':layout', 'The layout is ambiguous. Confirm long rows, a wide table, or a row-to-column matrix and its key/value fields.', f'{profile["name"]}: {", ".join(hs)}'))
                    continue
                sources.append(source)
                unused = [h for h in hs if h not in used]
                if unused:
                    problems.append(question(prefix + ':unused', 'What should happen to unmapped columns: ' + ', '.join(unused) + '? An explicit exclusion reason or mapping is required.', profile['name']))
    if mode == 'direct' and direct_candidates:
        multi_join = len(direct_candidates) > 1
        all_covered = {f for dc in direct_candidates for f in (list(dc['field_map'].keys()) + ([dc['list_field']] if dc['list_field'] else []))}
        if not multi_join:
            dc = direct_candidates[0]
            for field in headers:
                if field not in all_covered:
                    problems.append(question(dc['prefix'] + ':' + field, f'Which source column maps to {field}? Found {dc["match_counts"].get(field, 0)} matching headings.', dc['profile_name']))
            used_cols = set(dc['field_map'].values()) | set(dc['list_cols'])
            unused = [h for h in dc['hs'] if h not in used_cols]
            ignored_cols = {}
            if is_auto_template and len(all_covered) == len(headers) and unused:
                ignored_cols = {u: 'Auto-excluded extra column not in destination template' for u in unused}
                unused = []
            sources.append(SourcePlan(source_id=dc['profile_id'], sheet=dc['sheet_name'], header_row=dc['header_row'], layout='direct', key=key, period=dc['list_field'], columns=dc['list_cols'], field_map=dc['field_map'], ignored_columns=ignored_cols))
            if unused:
                problems.append(question(dc['prefix'] + ':unused', 'What should happen to unmapped columns: ' + ', '.join(unused) + '? An explicit exclusion reason or mapping is required.', dc['profile_name']))
        else:
            for field in headers:
                if field not in all_covered:
                    problems.append(question('multi:' + field, f'Which source column maps to {field}? Found 0 matching headings across uploaded sheets.', template['name']))
            for dc in direct_candidates:
                used_cols = set(dc['field_map'].values()) | set(dc['list_cols'])
                ignored_cols = {h: 'Excluded auxiliary column in multi-sheet join' for h in dc['hs'] if h not in used_cols}
                sources.append(SourcePlan(source_id=dc['profile_id'], sheet=dc['sheet_name'], header_row=dc['header_row'], layout='direct', key=key, period=dc['list_field'], columns=dc['list_cols'], field_map=dc['field_map'], ignored_columns=ignored_cols))
    for item in unknown_for_job(job):
        problems.append(question('rule-' + digest(item)[:12], 'Confirm how this workbook instruction should become a supported validation rule: ' + item['text'], item['ref'], kind='rule'))
    for profile in [template] + job['profiles']['sources']:
        for feature in profile['unsupported_features']:
            problems.append(question('feature-' + digest(feature)[:12], feature, profile['name'], kind='unsupported'))
    return MappingPlan(destination=destination, sources=sources, rules=[Rule(**r) for r in rules_for_job(job)], questions=problems, unsupported_rules=[item['ref'] for item in unknown_for_job(job)])


def rebind_knowledge(job, item):
    profiles = job['profiles']['sources']
    by_signature = {}
    for profile in profiles:
        signature = digest([s['layout_fingerprint'] for s in profile['sheets']])
        by_signature.setdefault(signature, []).append(profile['id'])
    bindings = {}
    for old_id, signature in item['bindings'].items():
        candidates = by_signature.get(signature, [])
        if not candidates:
            return None
        bindings[old_id] = candidates.pop(0)
    payload = copy.deepcopy(item['plan'])
    for source in payload['sources']:
        if source['source_id'] not in bindings:
            return None
        source['source_id'] = bindings[source['source_id']]
    for rule in payload['rules']:
        for old_id, new_id in bindings.items():
            if rule['id'].startswith(old_id + ':'):
                rule['id'] = new_id + rule['id'][len(old_id):]
            if rule['evidence'].startswith(old_id + ':'):
                rule['evidence'] = new_id + rule['evidence'][len(old_id):]
    return MappingPlan.model_validate(payload)


def decimal_value(value, config=None):
    config = config or {}
    if isinstance(value, bool) or value is None:
        raise ValueError('A numeric value is missing or has the wrong type.')
    if isinstance(value, (int, float, Decimal)):
        result = Decimal(str(value))
    else:
        text = str(value).strip()
        thousand = config.get('thousands_separator', '')
        if '.' in text and not config.get('decimal_separator') and thousand != '.':
            raise ValueError('Numeric text with a decimal/grouping separator needs an explicit locale rule; dot is not silently assumed.')
        if thousand and thousand == config.get('decimal_separator'):
            raise ValueError('Decimal and thousands separators conflict.')
        decimal = config.get('decimal_separator', '.')
        if thousand:
            parts = text.lstrip('+-').split(decimal)
            if len(parts) > 2 or (len(parts) == 2 and thousand in parts[1]):
                raise ValueError('Thousands grouping conflicts with the declared decimal locale.')
            groups = parts[0].split(thousand)
            if len(groups) > 1 and (not 1 <= len(groups[0]) <= 3 or any(len(g) != 3 or not g.isdigit() for g in groups[1:])):
                raise ValueError('Invalid thousands grouping for the declared locale.')
            text = text.replace(thousand, '')
        if decimal != '.':
            text = text.replace(decimal, '.')
        # Never silently choose a date/number locale or strip a currency symbol.
        if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)', text):
            raise ValueError('Numeric text requires explicit decimal/thousands rules; currency symbols are not guessed.')
        result = Decimal(text)
    if not result.is_finite():
        raise ValueError('Non-finite numbers are not supported.')
    return result


def numeric_or_original(value):
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return decimal_value(value)
    return value


def format_field_value(target_field: str, value, numeric_fields: set[str], config=None):
    if target_field in numeric_fields:
        return decimal_value(value, config)
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    nf = norm(target_field)
    if nf in {'grade', 'gr', 'class grade', 'student grade', 'learner grade', 'year level'}:
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            d = Decimal(str(value))
            s = str(int(d)) if d == d.to_integral_value() else str(value).strip()
        else:
            s = str(value).strip()
        m = re.fullmatch(r'(?:grade|gr\.?|g)?\s*(\d{1,2}[a-zA-Z]?)', s, re.I)
        if m:
            return f'Grade {m.group(1)}'
        return s
    if nf in {'date of birth', 'dob', 'birth date', 'birthdate', 'date_of_birth', 'learner dob'}:
        if hasattr(value, 'strftime'):
            return value.strftime('%Y-%m-%d')
        s = str(value).strip()
        m1 = re.match(r'^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})', s)
        if m1:
            return f'{int(m1[1]):04d}-{int(m1[2]):02d}-{int(m1[3]):02d}'
        m2 = re.match(r'^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})', s)
        if m2:
            return f'{int(m2[3]):04d}-{int(m2[2]):02d}-{int(m2[1]):02d}'
        return s
    if any(w in nf for w in ('phone', 'cell', 'mobile', 'telephone', 'contact number')):
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            d = Decimal(str(value))
            s = str(int(d)) if d == d.to_integral_value() else str(value).strip()
        else:
            s = str(value).strip()
        if s.startswith('+27'):
            digits = '0' + re.sub(r'\D', '', s[3:])
        else:
            digits = re.sub(r'[\s()-]', '', s)
        if re.fullmatch(r'[1-9]\d{8}', digits):
            return '0' + digits
        return digits
    if nf in {'subjects', 'subject', 'enrolled subjects', 'subject list', 'courses', 'modules'} and isinstance(value, str):
        parts = [p.strip() for p in re.split(r'[;,|]+', value) if p.strip()]
        return '; '.join(parts)
    if any(w in nf for w in ('sa id', 'id number', 'identity number', 'id no')):
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            d = Decimal(str(value))
            return str(int(d)) if d == d.to_integral_value() else str(value).strip()
        return str(value).strip()
    return numeric_or_original(value)


def sa_id_luhn_check_digit(first_12: str) -> int:
    odd_sum = sum(int(first_12[i]) for i in (0, 2, 4, 6, 8, 10))
    even_sum = sum((int(first_12[i]) * 2) if (int(first_12[i]) * 2) < 10 else (int(first_12[i]) * 2 - 9) for i in (1, 3, 5, 7, 9, 11))
    return (10 - ((odd_sum + even_sum) % 10)) % 10


def validate_and_repair_sa_id(value) -> dict:
    if value is None or (isinstance(value, str) and not value.strip()):
        return {'id_str': '', 'dob': None, 'gender': None, 'citizenship': None, 'valid_luhn': False, 'repaired': False, 'status': 'Missing ID'}
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        d = Decimal(str(value))
        raw = str(int(d)) if d == d.to_integral_value() else str(value).strip()
    else:
        raw = str(value).strip()
    # Check if alphanumeric foreign passport / study permit
    if re.search(r'[A-Za-z]', raw) and not re.match(r'^\d+(\.\d+)?[eE]\+\d+$', raw):
        return {'id_str': raw, 'dob': None, 'gender': None, 'citizenship': 'Foreign / Permit', 'valid_luhn': False, 'repaired': False, 'status': 'Foreign / Passport / Permit ID'}
    digits = re.sub(r'\D', '', raw)
    repaired = False
    if len(digits) in {11, 12}:
        padded = digits.zfill(13)
        mm_p, dd_p = int(padded[2:4]), int(padded[4:6])
        if 1 <= mm_p <= 12 and 1 <= dd_p <= 31:
            digits = padded
            repaired = True
    if len(digits) != 13:
        return {'id_str': raw, 'dob': None, 'gender': None, 'citizenship': None, 'valid_luhn': False, 'repaired': False, 'status': f'Invalid Length ({len(digits)} digits)'}
    yy, mm, dd = int(digits[:2]), int(digits[2:4]), int(digits[4:6])
    if not (1 <= mm <= 12 and 1 <= dd <= 31):
        return {'id_str': digits, 'dob': None, 'gender': None, 'citizenship': None, 'valid_luhn': False, 'repaired': repaired, 'status': 'Invalid Calendar Date in SA ID'}
    century = 2000 if yy <= 26 else 1900
    dob = f'{century + yy:04d}-{mm:02d}-{dd:02d}'
    gender = 'Male' if int(digits[6:10]) >= 5000 else 'Female'
    cit = 'SA Citizen' if digits[10] == '0' else ('Permanent Resident' if digits[10] == '1' else 'Other')
    exp_z = sa_id_luhn_check_digit(digits[:12])
    luhn_ok = int(digits[12]) == exp_z
    if repaired and luhn_ok:
        status = f'Repaired SA ID (Leading 0 Restored · {cit})'
    elif repaired and not luhn_ok:
        status = f'Repaired Leading 0 (Checksum Warning — expected {exp_z})'
    elif luhn_ok:
        status = f'Valid SA ID ({cit})'
    else:
        status = f'Invalid SA ID (Checksum Failed — expected {exp_z})'
    return {'id_str': digits, 'dob': dob, 'gender': gender, 'citizenship': cit, 'valid_luhn': luhn_ok, 'repaired': repaired, 'status': status}


def percentage_to_caps_level(value) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    s = str(value).strip().rstrip('%').strip()
    if not s or norm(s) in {'-', '—', '--', 'n/a', 'na', 'none', 'null', 'absent', 'exempt'}:
        return None
    try:
        pct = float(s)
    except ValueError:
        return None
    if pct < 0:
        return None
    if pct <= 29.99:
        return 1
    if pct <= 39.99:
        return 2
    if pct <= 49.99:
        return 3
    if pct <= 59.99:
        return 4
    if pct <= 69.99:
        return 5
    if pct <= 79.99:
        return 6
    return 7


def is_present_subject_mark(value) -> bool:
    if value is None or isinstance(value, bool):
        return False
    s = str(value).strip()
    if not s:
        return False
    if norm(s) in {'-', '—', '--', 'n/a', 'na', 'none', 'null', 'nil', 'no', 'false', 'not taken', 'exempt', 'absent'}:
        return False
    return True


def normalize_sources(job, plan):
    files = {f['id']: Path(f['path']) for f in job['files']['sources']}
    records, questions, exclusions = [], [], []
    expected = 0
    config = job['profiles']['template']['config']
    numeric_fields = {field for rule in plan.rules if rule.kind in {'nonnegative', 'numeric', 'integer', 'min', 'max', 'sum'} for field in rule.fields}
    numeric_fields |= {field.strip() for rule in plan.rules if rule.kind == 'sum' for field in rule.value.split(',') if field.strip()}
    ignored_labels = {'subtotal', 'total', 'grand total', 'sum', 'average', 'notes', 'source:'} | {norm(v) for v in config.get('ignore_row_labels', '').split('|') if v}
    seen_regions = set()
    for source in plan.sources:
        region = (source.source_id, source.sheet)
        if region in seen_regions:
            questions.append(question('duplicate-region-' + digest(region)[:8], 'The same source region is mapped twice. Remove the duplicate mapping.', str(region), kind='mapping'))
            continue
        seen_regions.add(region)
        if source.source_id not in files:
            questions.append(question('unknown-source', 'Mapping references a source that is not in this job.', source.source_id))
            continue
        raw_book = book_rows(files[source.source_id])
        cached_book = book_rows(files[source.source_id], data_only=True)
        if source.sheet not in raw_book:
            questions.append(question('unknown-sheet', 'Mapping references a missing source sheet.', source.sheet))
            continue
        raw_rows, rows = raw_book[source.sheet], cached_book[source.sheet]
        source_profile = next(p for p in job['profiles']['sources'] if p['id'] == source.source_id)
        sheet_profile = next(s for s in source_profile['sheets'] if s['name'] == source.sheet)
        formula_coordinates = set(sheet_profile['formula_cells'])
        if source.layout == 'cells':
            expected += 1
            if plan.destination.mode != 'direct':
                questions.append(question('form-destination', 'Row-label records require direct output mode in this version.', source.sheet, kind='unsupported'))
                continue
            record, cell_questions, cell_exclusions = normalise_cell_form(source, job, raw_rows, rows, numeric_fields, config)
            records.append(record)
            questions.extend(cell_questions)
            exclusions.extend(cell_exclusions)
            continue
        if source.header_row > len(rows):
            questions.append(question('header-range', 'Header row is outside the source sheet.', source.sheet))
            continue
        headers = [str(v or '').strip() for v in raw_rows[source.header_row - 1]]
        while headers and not headers[-1]:
            headers.pop()
        if len(set(norm(h) for h in headers)) != len(headers) or any(not h for h in headers):
            questions.append(question('bad-source-headers', 'Blank or repeated source headings must be resolved.', source.sheet))
            continue
        index = {h: i for i, h in enumerate(headers)}
        needed = (set(source.field_map.values()) | set(source.columns)) if source.layout == 'direct' else set(x for x in [source.key, source.period, source.value] + source.columns if x)
        absent = needed - set(headers)
        unaccounted = set(headers) - needed - set(source.ignored_columns)
        if absent:
            questions.append(question('missing-column-' + digest(region)[:8], 'Mapped source headings do not exist: ' + ', '.join(sorted(absent)), source.sheet))
            continue
        if source.layout == 'direct' and not job.get('plan_confirmation'):
            for target_field, column in source.field_map.items():
                if job['profiles']['template']['aliases'].get(target_field) and norm(column) not in aliases_for(job['profiles']['template'], target_field):
                    questions.append(question('alias-' + digest([region, target_field])[:10], f'The proposed {column} → {target_field} mapping differs from the template-defined aliases. Confirm its meaning before release.', source.sheet, kind='mapping'))
        if unaccounted:
            questions.append(question('unmapped-column-' + digest(region)[:8], 'Unmapped columns need explicit disposition: ' + ', '.join(sorted(unaccounted)), source.sheet))
        for column, reason in source.ignored_columns.items():
            if not reason.strip():
                questions.append(question('no-exclusion-reason', 'An excluded column has no reason.', column))
            exclusions.append({'source_id': source.source_id, 'sheet': source.sheet, 'column': column, 'reason': reason})
        role_label = source.field_map.get(source.key) or source.key or source.period or (headers[0] if headers else '')
        if role_label not in index and headers:
            role_label = headers[0]
        prospective_rows = [row for row in raw_rows[source.header_row:] if any(v is not None and str(v).strip() for v in row) and [norm(v) for v in row] != [norm(h) for h in headers] and not (role_label in index and norm(row[index[role_label]] if index[role_label] < len(row) else None) in ignored_labels)]
        projected = len(prospective_rows) * (len(source.columns) if source.layout in {'matrix', 'wide'} else 1)
        if expected + projected > MAX_NORMALIZED_RECORDS:
            expected += projected
            questions.append(question('record-limit-' + digest(region)[:8], f'This job exceeds the {MAX_NORMALIZED_RECORDS:,} normalized-record safety limit. Split the job or approve a measured capacity change.', source.sheet, kind='unsupported'))
            continue
        for number, prefix in enumerate(raw_rows[:source.header_row - 1], 1):
            nonempty = [v for v in prefix if v is not None and str(v).strip()]
            if not nonempty:
                continue
            exclusions.append({'source_id': source.source_id, 'sheet': source.sheet, 'row': number, 'reason': 'Pre-header title/metadata region; not part of the selected data table.'})
        verified_formulas, visiting_formulas = {}, set()
        def source_lookup(coordinate):
            col, row_number, _, _ = range_boundaries(coordinate + ':' + coordinate)
            offset_row, offset_col = row_number - 1, col - 1
            if offset_row >= len(raw_rows) or offset_col >= len(raw_rows[offset_row]):
                raise ValueError('Source formula references a missing cell.')
            original = raw_rows[offset_row][offset_col]
            cached = rows[offset_row][offset_col]
            if coordinate in formula_coordinates and isinstance(original, str) and original.startswith('='):
                if coordinate in verified_formulas:
                    return verified_formulas[coordinate]
                if coordinate in visiting_formulas:
                    raise ValueError('Circular source formula needs correction.')
                if cached is None:
                    raise ValueError(f'{coordinate} has a formula without a verified cached result. Recalculate the source workbook first.')
                visiting_formulas.add(coordinate)
                computed = compute_formula(original, source_lookup)
                visiting_formulas.remove(coordinate)
                if computed != decimal_value(cached):
                    raise ValueError(f'Cached source formula result disagrees at {coordinate}.')
                verified_formulas[coordinate] = computed
                return computed
            return cached
        for number, row in enumerate(rows[source.header_row:], source.header_row + 1):
            raw = raw_rows[number - 1]
            if all(v is None or str(v).strip() == '' for v in raw):
                continue
            if [norm(v) for v in raw] == [norm(v) for v in headers]:
                exclusions.append({'source_id': source.source_id, 'sheet': source.sheet, 'row': number, 'reason': 'Repeated header, identical to the selected header row.'})
                continue
            label_col = source.field_map.get(source.key) or source.key or source.period or (headers[0] if headers else '')
            if label_col not in index and headers:
                label_col = headers[0]
            label = row[index[label_col]] if label_col in index and index[label_col] < len(row) else None
            if norm(label) in ignored_labels:
                exclusions.append({'source_id': source.source_id, 'sheet': source.sheet, 'row': number, 'reason': 'Explicit ignore_row_labels workbook rule.'})
                continue
            def get(column):
                offset = index[column]
                value = row[offset] if offset < len(row) else None
                original = raw[offset] if offset < len(raw) else None
                coordinate = f'{get_column_letter(offset + 1)}{number}'
                if coordinate in formula_coordinates and isinstance(original, str) and original.startswith('='):
                    try:
                        value = source_lookup(coordinate)
                    except (ValueError, SyntaxError, InvalidOperation, ZeroDivisionError) as exc:
                        raise ValueError(f'Source formula {coordinate} cannot be verified: {exc}') from exc
                return value, coordinate
            try:
                if source.layout == 'direct':
                    expected += 1
                    data, refs = {}, []
                    for target_field, column in source.field_map.items():
                        value, cell = get(column)
                        data[target_field] = format_field_value(target_field, value, numeric_fields, config)
                        refs.append({'source_id': source.source_id, 'sheet': source.sheet, 'cell': cell, 'target_field': target_field})
                    if source.period and source.columns:
                        taken = []
                        for col in source.columns:
                            val, cell = get(col)
                            if is_present_subject_mark(val):
                                clean_subj = re.sub(r'\s*(?:\(%?\)|%|mark|percent|percentage)\s*$', '', col, flags=re.I).strip() or col.strip()
                                taken.append(clean_subj)
                                refs.append({'source_id': source.source_id, 'sheet': source.sheet, 'cell': cell, 'target_field': source.period})
                        data[source.period] = '; '.join(taken)
                    records.append({'id': f'{source.source_id}:{source.sheet}:{number}', 'data': data, 'refs': refs, 'operation': 'horizontal subject-percentage matrix + column mapping' if source.period else 'column mapping'})
                elif source.layout == 'long':
                    expected += 1
                    entity, key_cell = get(source.key)
                    category, period_cell = get(source.period)
                    value, value_cell = get(source.value)
                    if entity is None or category is None:
                        raise ValueError('A record key or category is missing.')
                    records.append({'id': f'{source.source_id}:{source.sheet}:{number}', 'entity': str(entity).strip(), 'category': str(category).strip(), 'value': decimal_value(value, config), 'refs': [{'source_id': source.source_id, 'sheet': source.sheet, 'cells': [key_cell, period_cell, value_cell]}]})
                else:
                    for column in source.columns:
                        coord = f'{get_column_letter(index[column] + 1)}{number}'
                        if coord in source.ignored_cells:
                            exclusions.append({'source_id': source.source_id, 'sheet': source.sheet, 'cell': coord, 'reason': source.ignored_cells[coord]})
                            continue
                        expected += 1
                        if source.layout == 'wide':
                            entity, key_cell = get(source.key)
                            category, period_cell = column, f'{get_column_letter(index[column] + 1)}{source.header_row}'
                        else:
                            entity, key_cell = column, f'{get_column_letter(index[column] + 1)}{source.header_row}'
                            category, period_cell = get(source.period)
                        try:
                            value, value_cell = get(column)
                            if entity is None or category is None:
                                raise ValueError('A record key or category is missing.')
                            records.append({'id': f'{source.source_id}:{source.sheet}:{number}:{index[column]}', 'entity': str(entity).strip(), 'category': str(category).strip(), 'value': decimal_value(value, config), 'refs': [{'source_id': source.source_id, 'sheet': source.sheet, 'cells': [key_cell, period_cell, value_cell]}]})
                        except (ValueError, InvalidOperation) as exc:
                            questions.append(question('cell-' + digest([region, number, column])[:10], str(exc), f'{source.sheet}!{get_column_letter(index[column] + 1)}{number}', kind='data'))
            except (ValueError, InvalidOperation, IndexError) as exc:
                questions.append(question('row-' + digest([region, number])[:10], str(exc), f'{source.sheet}!{number}', kind='data'))
    available = {(p['id'], s['name']) for p in job['profiles']['sources'] for s in p['sheets']}
    for region in sorted(available - seen_regions):
        questions.append(question('unmapped-sheet-' + digest(region)[:10], 'A source sheet has no mapping or approved scope disposition.', str(region)))
    return records, expected, questions, exclusions


def build_rows(records, destination, rules):
    questions, provenance = [], []
    categories = {norm(x): x for x in destination.categories}
    rows = []
    fields_count = len(destination.categories) + 1 + bool(destination.total) if destination.mode == 'pivot' else (3 if destination.mode == 'long' else max((len(r.get('data', {})) for r in records), default=1))
    max_rows = min(MAX_OUTPUT_ROWS, MAX_OUTPUT_CELLS // max(1, fields_count))
    if destination.mode == 'direct':
        source_regions = {r['id'].rsplit(':', 1)[0] for r in records}
        if len(source_regions) > 1:
            by_key = {}
            by_name = {}
            for record in records:
                data = record['data']
                k_val = norm(data.get(destination.key)) if destination.key else ''
                fn_ln = (norm(data.get('First Name') or data.get('first_name')), norm(data.get('Last Name') or data.get('last_name') or data.get('Surname')))
                reg = record['id'].rsplit(':', 1)[0]
                target_idx = None
                if k_val and k_val in by_key and by_key[k_val][1] != reg:
                    target_idx = by_key[k_val][0]
                elif fn_ln[0] and fn_ln[1] and fn_ln in by_name and by_name[fn_ln][1] != reg:
                    target_idx = by_name[fn_ln][0]
                if target_idx is not None:
                    merged_row = rows[target_idx]
                    for f, v in data.items():
                        if merged_row.get(f) in (None, '') and v not in (None, ''):
                            merged_row[f] = v
                    provenance.append({'record_id': record['id'], 'target_row': destination.header_row + target_idx + 1, 'target_fields': list(data), 'sources': record['refs'], 'operation': 'multi-sheet relational join + ' + record.get('operation', 'column mapping')})
                else:
                    if len(rows) >= max_rows:
                        questions.append(question('output-limit', f'Output exceeds the {max_rows:,}-row capacity for this template. No truncated workbook will be released.', kind='unsupported'))
                        break
                    new_idx = len(rows)
                    row_copy = dict(data)
                    rows.append(row_copy)
                    if k_val:
                        by_key[k_val] = (new_idx, reg)
                    if fn_ln[0] and fn_ln[1]:
                        by_name[fn_ln] = (new_idx, reg)
                    provenance.append({'record_id': record['id'], 'target_row': destination.header_row + new_idx + 1, 'target_fields': list(data), 'sources': record['refs'], 'operation': record.get('operation', 'column mapping')})
            return rows, provenance, questions
        for index, record in enumerate(records):
            if index >= max_rows:
                questions.append(question('output-limit', f'Output exceeds the {max_rows:,}-row capacity for this template. No truncated workbook will be released.', kind='unsupported'))
                break
            rows.append(record['data'])
            provenance.append({'record_id': record['id'], 'target_row': destination.header_row + index + 1, 'target_fields': list(record['data']), 'sources': record['refs'], 'operation': record.get('operation', 'column mapping')})
        return rows, provenance, questions
    if destination.mode == 'long':
        for index, record in enumerate(records):
            if index >= max_rows:
                questions.append(question('output-limit', f'Output exceeds the {max_rows:,}-row capacity for this template. No truncated workbook will be released.', kind='unsupported'))
                break
            category = categories.get(norm(record['category']))
            if category is None:
                questions.append(question('category-' + digest(record['id'])[:10], f'Category {record["category"]} is not in the destination rules.', record['id'], kind='data'))
                category = record['category']
            rows.append({destination.key: record['entity'], destination.period: category, destination.value: record['value']})
            provenance.append({'record_id': record['id'], 'target_row': destination.header_row + index + 1, 'target_fields': [destination.key, destination.period, destination.value], 'sources': record['refs'], 'operation': 'unpivot / long rows'})
        return rows, provenance, questions
    groups = OrderedDict()
    aggregate_fields = {field for r in rules if r.kind == 'aggregate' and r.value == 'sum' for field in r.fields}
    for record in records:
        if record['entity'] not in groups and len(groups) >= max_rows:
            questions.append(question('output-limit', f'Output exceeds the {max_rows:,}-row capacity for this template. No truncated workbook will be released.', kind='unsupported'))
            continue
        group = groups.setdefault(record['entity'], {destination.key: record['entity'], **{category: None for category in destination.categories}})
        category = categories.get(norm(record['category']))
        if category is None:
            questions.append(question('category-' + digest(record['id'])[:10], f'Category {record["category"]} does not match a destination column.', record['id'], kind='data'))
            continue
        if group[category] is not None and category not in aggregate_fields:
            questions.append(question('duplicate-' + digest([record['entity'], category])[:10], f'Multiple values map to {record["entity"]} / {category}. Confirm aggregation or correct the source.', record['id'], kind='data'))
        elif group[category] is None:
            group[category] = record['value']
        else:
            group[category] += record['value']
        provenance.append({'record_id': record['id'], 'target_row': destination.header_row + list(groups).index(record['entity']) + 1, 'target_fields': [destination.key, category], 'sources': record['refs'], 'operation': 'rows to columns / pivot'})
    for entity, group in groups.items():
        for category in destination.categories:
            if group[category] is None:
                questions.append(question('missing-' + digest([entity, category])[:10], f'No source value for {entity} / {category}. Supply it or an explicitly authorised rule; zero is not assumed.', destination.sheet, kind='data'))
        if destination.total:
            group[destination.total] = sum((group[c] for c in destination.categories), Decimal(0)) if all(group[c] is not None for c in destination.categories) else None
        rows.append(group)
    return rows, provenance, questions


def evaluate_rules(rows, headers, rules):
    checks = []
    for rule in rules:
        missing = set(rule.fields) - set(headers)
        errors = []
        if missing:
            errors.append('Rule refers to unknown output fields: ' + ', '.join(sorted(missing)))
        elif rule.kind == 'aggregate':
            if rule.value != 'sum':
                errors.append('Only explicitly authorised sum aggregation is supported.')
        elif rule.kind == 'unique':
            seen = set()
            for index, row in enumerate(rows, 1):
                key = tuple(canonical(row.get(f)) for f in rule.fields)
                if key in seen:
                    errors.append(f'Duplicate key at output record {index}.')
                seen.add(key)
        else:
            for index, row in enumerate(rows, 1):
                for field in rule.fields:
                    value = row.get(field)
                    try:
                        if rule.kind == 'required':
                            if value is None or (isinstance(value, str) and not value.strip()):
                                errors.append(f'{field} is missing at record {index}.')
                        elif rule.kind == 'allowed':
                            if str(public_value(value)) not in rule.value.split('|'):
                                errors.append(f'{field} has a value outside the allowed list at record {index}.')
                        elif rule.kind == 'sum':
                            components = [x.strip() for x in rule.value.split(',') if x.strip()]
                            if not components or any(c not in headers for c in components):
                                errors.append('Sum rule has missing or unknown component fields.')
                            elif decimal_value(value) != sum((decimal_value(row.get(c)) for c in components), Decimal(0)):
                                errors.append(f'{field} does not reconcile at record {index}.')
                        else:
                            number = decimal_value(value)
                            if rule.kind == 'nonnegative' and number < 0:
                                errors.append(f'{field} is negative at record {index}.')
                            elif rule.kind == 'min' and number < Decimal(rule.value):
                                errors.append(f'{field} is below {rule.value} at record {index}.')
                            elif rule.kind == 'max' and number > Decimal(rule.value):
                                errors.append(f'{field} is above {rule.value} at record {index}.')
                            elif rule.kind == 'integer' and number != number.to_integral_value():
                                errors.append(f'{field} is not an integer at record {index}.')
                    except (ValueError, InvalidOperation, TypeError):
                        errors.append(f'{field} cannot be checked as required at record {index}.')
        checks.append(Check(id='rule-' + rule.id, name=f'{rule.id} · {rule.kind}', passed=not errors, detail='; '.join(errors[:3]) if errors else f'Passed for {len(rows)} output records. Source: {rule.evidence}'))
    return checks


def compute_formula(formula, lookup):
    expression = formula.lstrip('=').replace('$', '').strip()
    match = re.fullmatch(r'SUM\(([A-Z]+\d+):([A-Z]+\d+)\)', expression, re.I)
    if match:
        a, b, c, d = range_boundaries(match[1].upper() + ':' + match[2].upper())
        return sum((decimal_value(lookup(f'{get_column_letter(col)}{row}')) for row in range(b, d + 1) for col in range(a, c + 1)), Decimal(0))
    tree = ast.parse(expression, mode='eval')
    def visit(node):
        if isinstance(node, ast.Expression):
            return visit(node.body)
        if isinstance(node, ast.Name) and re.fullmatch(r'[A-Z]+[1-9]\d*', node.id):
            return decimal_value(lookup(node.id))
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) and not isinstance(node.value, bool):
            return Decimal(str(node.value))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            return visit(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
        if isinstance(node, ast.BinOp):
            left, right = visit(node.left), visit(node.right)
            if isinstance(node.op, ast.Add): return left + right
            if isinstance(node.op, ast.Sub): return left - right
            if isinstance(node.op, ast.Mult): return left * right
            if isinstance(node.op, ast.Div): return left / right
        raise ValueError('Formula uses an unsupported operation. No arbitrary formula code is executed.')
    return visit(tree)


def render_output(job, plan, rows, path, headers):
    wb = load_workbook(job['files']['template']['path'], keep_links=False)
    ws = wb[plan.destination.sheet]
    start = plan.destination.header_row + 1
    patterns = {}
    for col, field in enumerate(headers, 1):
        cell = ws.cell(start, col)
        patterns[field] = {'style': copy.copy(cell._style), 'formula': cell.value if cell.data_type == 'f' else None}
    if ws.max_row >= start:
        ws.delete_rows(start, ws.max_row - start + 1)
    values, formulas = {}, {}
    for index, row in enumerate(rows, start):
        for col, field in enumerate(headers, 1):
            cell = ws.cell(index, col)
            cell._style = copy.copy(patterns[field]['style'])
            value = row.get(field)
            if isinstance(value, Decimal) and len(value.normalize().as_tuple().digits) > 15:
                raise ValueError('Numeric precision exceeds Excel’s 15-significant-digit boundary. Use an explicitly typed text identifier or a different verified representation.')
            if patterns[field]['formula']:
                cell.value = Translator(patterns[field]['formula'], origin=f'{get_column_letter(col)}{start}').translate_formula(cell.coordinate)
                formulas[cell.coordinate] = cell.value
                values[cell.coordinate] = value
            else:
                cell.value = value
                if isinstance(value, str) and (value.startswith(('=', '+', '-', '@')) or re.fullmatch(r'0\d+|\d{13}', value)):
                    cell.data_type = 's'  # Source content is data, not executable Excel formulas or stripped leading-zero / 13-digit IDs.
                values[cell.coordinate] = value
    computed, visiting = {}, set()
    def lookup(coordinate):
        if coordinate in computed:
            return computed[coordinate]
        if coordinate in formulas:
            if coordinate in visiting:
                raise ValueError('Circular template formula detected.')
            visiting.add(coordinate)
            answer = compute_formula(formulas[coordinate], lookup)
            visiting.remove(coordinate)
            expected = values.get(coordinate)
            if expected is None or answer != decimal_value(expected):
                raise ValueError(f'Template formula {coordinate} disagrees with the transformed value.')
            computed[coordinate] = answer
            return answer
        if coordinate not in values:
            raise ValueError(f'Formula references an unsupported or missing cell: {coordinate}.')
        return values[coordinate]
    for coordinate in formulas:
        lookup(coordinate)
    ws.auto_filter.ref = f'A{plan.destination.header_row}:{get_column_letter(len(headers))}{max(start, start + len(rows) - 1)}'
    ws.freeze_panes = f'A{start}'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    sheet_index = wb.worksheets.index(ws) + 1
    wb.close()
    if computed:
        member = f'xl/worksheets/sheet{sheet_index}.xml'
        namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        root = ET.fromstring(entries[member])
        for cell in root.iter(namespace + 'c'):
            if cell.attrib.get('r') in computed:
                v = cell.find(namespace + 'v')
                if v is None:
                    v = ET.SubElement(cell, namespace + 'v')
                v.text = str(computed[cell.attrib['r']])
        entries[member] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
        temporary = path.with_suffix('.cached.xlsx')
        with zipfile.ZipFile(temporary, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        temporary.replace(path)
    exported = book_rows(path, data_only=True)[plan.destination.sheet]
    for offset, row in enumerate(rows, start - 1):
        for col, field in enumerate(headers):
            got = exported[offset][col] if col < len(exported[offset]) else None
            expected = row.get(field)
            if isinstance(expected, Decimal):
                if decimal_value(got) != expected:
                    raise ValueError('Exported numeric value lost precision or differs from the validated value.')
            elif got != expected:
                raise ValueError('Reopened workbook differs from the validated output.')
    return hashlib.sha256(path.read_bytes()).hexdigest()


def transform(job, plan, output_path, knowledge_conflict=False, use_workbook_rules=False):
    template = job['profiles']['template']
    config = template['config']
    checks = []
    questions = list(plan.questions)
    def check(identifier, name, passed, detail):
        checks.append(Check(id=identifier, name=name, passed=passed, detail=detail))
    input_files = [job['files']['template']] + job['files']['sources']
    fingerprints = {p['id']: p['sha256'] for p in [template] + job['profiles']['sources']}
    intact = all(Path(f['path']).exists() and hashlib.sha256(Path(f['path']).read_bytes()).hexdigest() == fingerprints[f['id']] for f in input_files)
    check('input-integrity', 'Original input fingerprints unchanged', intact, 'Current original-file bytes must match the inspected manifest.')
    target_sheets = book_rows(Path(job['files']['template']['path']))
    destination_valid = plan.destination.sheet in target_sheets and plan.destination.header_row <= len(target_sheets.get(plan.destination.sheet, []))
    check('target-region', 'Destination region exists', destination_valid, 'Mapping must reference a real destination sheet and header row.')
    if not destination_valid:
        questions.append(question('target-region', 'Choose a real destination region.', kind='mapping'))
        headers = []
    else:
        headers = [str(v or '').strip() for v in target_sheets[plan.destination.sheet][plan.destination.header_row - 1]]
        while headers and not headers[-1]:
            headers.pop()
    shape_fields = ([plan.destination.key] + plan.destination.categories + ([plan.destination.total] if plan.destination.total else [])) if plan.destination.mode == 'pivot' else ([plan.destination.key, plan.destination.period, plan.destination.value] if plan.destination.mode == 'long' else list({k for source in plan.sources for k in (list(source.field_map) + ([source.period] if source.period else []))}))
    check('destination-fields', 'Every destination field is mapped', bool(headers) and len(headers) == len(set(norm(h) for h in headers)) and all(headers) and set(headers) == set(shape_fields), 'Output headers must be unique, non-empty and fully covered by the plan.')
    check('conflicts', 'Rule conflicts resolved', not knowledge_conflict or use_workbook_rules, 'Saved rules differ from this workbook: an explicit clarification is required.' if knowledge_conflict and not use_workbook_rules else 'No unresolved saved/workbook rule conflict.')
    if knowledge_conflict and not use_workbook_rules:
        questions.append(question('knowledge-conflict', 'Saved knowledge and current workbook rules conflict. Which should govern this template version?', 'Rules & Knowledge', kind='conflict', options=['Use current workbook rules for this version', 'Keep this job held']))
    known = rules_for_job(job)
    proposed = {canonical(r.model_dump()) for r in plan.rules}
    check('rule-integrity', 'Authoritative rules preserved', all(canonical(r) in proposed for r in known), 'Parsed workbook rules cannot be silently changed or dropped by a mapping plan.')
    unknown = unknown_for_job(job)
    is_auto_template = any(r.get('id', '').startswith('AUTO-') for r in template['rules'])
    manual_confirmation = bool(job.get('plan_confirmation')) or is_auto_template
    covered_refs = {r.evidence for r in plan.rules}
    uncompiled = [item for item in unknown if item['ref'] not in covered_refs or not manual_confirmation]
    check('rule-coverage', 'All instructions compiled or clarified', not uncompiled and not plan.unsupported_rules, 'Uncompiled instructions: ' + ', '.join(x['ref'] for x in uncompiled) if uncompiled else 'All discovered instructions have supported checks or recorded confirmation.')
    for item in uncompiled:
        questions.append(question('uncompiled-' + digest(item)[:10], 'This instruction still requires a supported rule and explicit clarification: ' + item['text'], item['ref'], kind='rule'))
    check('rule-exists', 'Authoritative quality rules are present', any(r['kind'] == 'required' for r in known) or (manual_confirmation and any(r.kind == 'required' for r in plan.rules)), 'At least one authoritative required-field rule is needed; no-rule jobs cannot be certified.')
    config_ok = (not config.get('mode') or config['mode'] == plan.destination.mode) and (not config.get('key') or config['key'] == plan.destination.key)
    check('config-integrity', 'Workbook configuration respected', config_ok, 'The destination mode and key must match explicit workbook instructions.')
    check('write-policy', 'Destination replacement authorised', config.get('write_mode') == 'replace', 'This version requires an explicit write_mode: replace instruction; append and multi-block output are not implemented.')
    features = [f for p in [template] + job['profiles']['sources'] for f in p['unsupported_features']]
    check('supported-features', 'Workbook features supported', not features, '; '.join(features) if features else 'No unsupported workbook feature was detected by the current scanner.')
    aggregate_rules = [r for r in plan.rules if r.kind == 'aggregate']
    if aggregate_rules:
        authorised = plan.destination.mode == 'pivot' and all(r.value == 'sum' and set(r.fields).issubset(set(plan.destination.categories)) for r in aggregate_rules)
        check('aggregation-scope', 'Aggregation has explicit field scope', authorised, 'Sum aggregation is supported only for the destination pivot categories explicitly named by the rule.')
        if not authorised:
            questions.append(question('aggregation-scope', 'Clarify the exact pivot fields authorised for aggregation; other aggregation shapes are not implemented.', kind='rule'))
    auto_reasons = {'Excluded auxiliary column in multi-sheet join', 'Auto-excluded summary column', 'Auto-excluded summary/extra column', 'Auto-excluded extra column not in destination template'}
    has_exclusions = any(
        s.ignored_cells or any(r not in auto_reasons for r in s.ignored_columns.values())
        for s in plan.sources
    )
    if has_exclusions:
        check('exclusion-authority', 'Exclusions have operator authority', manual_confirmation, 'AI alone cannot authorise dropping columns or cells. Confirm the exact scope explicitly.')
        if not manual_confirmation:
            questions.append(question('exclusion-authority', 'Confirm which columns/cells may be excluded and why before QMS can release this output.', kind='mapping'))
    if features:
        questions.append(question('unsupported-workbook-features', 'This workbook needs a verified operation before release: ' + '; '.join(features), kind='unsupported'))
    records, expected, data_questions, exclusions = normalize_sources(job, plan)
    questions += data_questions
    rows, provenance, row_questions = build_rows(records, plan.destination, plan.rules)
    questions += row_questions
    mapped_ids = {p['record_id'] for p in provenance}
    check('source-coverage', 'In-scope records fully accounted for', expected > 0 and len(records) == expected and len(mapped_ids) == expected, f'{len(mapped_ids)} mapped / {expected} in-scope records. Structural exclusions are listed in the report.')
    check('nonempty', 'Output contains records', bool(rows), f'{len(rows)} destination records generated.')
    if plan.destination.mode in {'pivot', 'long'}:
        total_in = sum((r['value'] for r in records), Decimal(0))
        if plan.destination.mode == 'pivot':
            values = [row.get(category) for row in rows for category in plan.destination.categories]
        else:
            values = [row.get(plan.destination.value) for row in rows]
        complete = all(value is not None for value in values)
        total_out = sum((decimal_value(v) for v in values), Decimal(0)) if complete else None
        check('reconciliation', 'Source values reconcile to output', complete and total_in == total_out, f'Input value sum: {total_in}; output value sum: {total_out if complete else "incomplete"}. Derived totals are not double-counted.')
    checks += evaluate_rules(rows, headers, plan.rules)
    for failed in checks:
        if not failed.passed and failed.id.startswith('rule-'):
            questions.append(question('check-' + failed.id, failed.detail, failed.name, kind='rule'))
    # Deduplicate identical clarification IDs without hiding distinct failed checks.
    questions = list({q.id: q for q in questions}.values())
    check('clarifications', 'No unanswered clarification', not questions, f'{len(questions)} questions require attention.' if questions else 'No unresolved mapping, data or rule question remains.')
    output_hash = ''
    if all(c.passed for c in checks):
        try:
            output_hash = render_output(job, plan, rows, output_path, headers)
            check('export-roundtrip', 'Export reopened and verified', True, 'Values match the validated result; supported template formulas and styles are preserved.')
        except Exception as exc:
            check('export-roundtrip', 'Export reopened and verified', False, str(exc))
            questions.append(question('export-validation', str(exc), plan.destination.sheet, kind='unsupported'))
            output_path.unlink(missing_ok=True)
    else:
        check('export-roundtrip', 'Export gated by QMS', False, 'Final workbook is not created or downloadable while a required check fails.')
        output_path.unlink(missing_ok=True)
    if questions:
        for gate in checks:
            if gate.id == 'clarifications':
                gate.passed = False
                gate.detail = f'{len({q.id for q in questions})} questions require attention.'
    released = all(c.passed for c in checks) and not questions
    report = Report(status='RELEASED' if released else 'HELD', checks=checks, questions=list({q.id: q for q in questions}.values()), rows=[{k: public_value(v) for k, v in row.items()} for row in rows[:200]], headers=headers, source_records=expected, mapped_records=len(mapped_ids), output_rows=len(rows), provenance=provenance, exclusions=exclusions, mapping_origin=job.get('mapping_origin', 'rule-informed layout resolver'), template_version=template['version'], trace_id=job['trace_id'], passed_checks=sum(c.passed for c in checks), total_checks=len(checks), output_sha256=output_hash)
    return report
