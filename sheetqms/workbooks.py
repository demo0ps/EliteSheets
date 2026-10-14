from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from .models import Rule
from .store import digest

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLS = 200
MAX_CELLS = 1_000_000
RULE_SHEETS = {'rules', 'instructions', 'mapping rules', 'business rules'}
KINDS = set(Rule.model_fields['kind'].annotation.__args__)
CONFIG_KEYS = {'mode', 'key', 'period', 'value', 'categories', 'total', 'header_row', 'output_sheet', 'write_mode', 'ignore_row_labels', 'decimal_separator', 'thousands_separator'}
PERIOD_TOKENS = {'jan', 'january', 'feb', 'february', 'mar', 'march', 'apr', 'april', 'may', 'jun', 'june', 'jul', 'july', 'aug', 'august', 'sep', 'sept', 'september', 'oct', 'october', 'nov', 'november', 'dec', 'december', 'q1', 'q2', 'q3', 'q4', 'w1', 'w2', 'w3', 'w4'}
BUILTIN_SYNONYMS = [
    ['site', 'branch', 'location', 'store', 'region', 'office', 'plant', 'depot', 'facility', 'warehouse', 'center', 'centre', 'entity', 'division', 'market'],
    ['item', 'item code', 'item_code', 'item no', 'item id', 'sku', 'part', 'part number', 'part no', 'product id', 'product code', 'code', 'id', 'material', 'article', 'catalog no', 'ref'],
    ['description', 'name', 'product', 'product name', 'item name', 'item description', 'material description', 'title', 'details', 'summary', 'label', 'particulars'],
    ['qty', 'quantity', 'units', 'unit count', 'count', 'ordered', 'shipped', 'delivered', 'volume', 'stock', 'on hand', 'pieces', 'pcs'],
    ['price', 'unit price', 'unit_price', 'cost', 'unit cost', 'rate', 'price each', 'list price'],
    ['total', 'amount', 'total amount', 'line total', 'ext price', 'extended price', 'revenue', 'sales', 'net amount', 'gross amount', 'subtotal', 'sum', 'value'],
    ['date', 'order date', 'invoice date', 'ship date', 'transaction date', 'day', 'created', 'timestamp'],
    ['period', 'month', 'quarter', 'year', 'interval', 'week', 'fiscal period'],
    ['status', 'state', 'stage', 'condition', 'order status'],
    ['category', 'type', 'group', 'department', 'dept', 'class', 'family', 'segment'],
    ['customer', 'client', 'buyer', 'account', 'customer name', 'client name'],
    ['supplier', 'vendor', 'manufacturer', 'partner', 'supplier name', 'vendor name'],
    ['notes', 'comments', 'remarks', 'memo', 'observation'],
    ['student id', 'learner id', 'accession number', 'accession no', 'admission no', 'admission number', 'lurits', 'lurits no', 'student no', 'student number', 'learner no', 'student_id', 'sa id', 'sa id number', 'id number', 'number', 'no', 'no.'],
    ['first name', 'firstname', 'learner first name', 'student first name', 'given name', 'forenames', 'first_name'],
    ['last name', 'lastname', 'surname', 'learner surname', 'student surname', 'family name', 'last_name'],
    ['date of birth', 'dob', 'birth date', 'birthdate', 'date_of_birth', 'learner dob'],
    ['grade', 'gr', 'class grade', 'student grade', 'learner grade', 'year level'],
    ['subjects', 'subject', 'enrolled subjects', 'subject list', 'courses', 'modules'],
    ['parent first name', 'parent firstname', 'guardian first name', 'parent_first_name', 'parent given name', 'parent 1 first name', 'parent 2 first name', 'father name', 'mother name'],
    ['parent last name', 'parent lastname', 'parent surname', 'guardian surname', 'guardian last name', 'parent_last_name', 'parent 1 last name', 'parent 2 last name', 'father surname', 'mother surname'],
    ['parent email', 'guardian email', 'email', 'email address', 'parent_email', 'contact email', 'parent 1 email', 'parent 2 email'],
    ['parent phone', 'parent cell', 'guardian phone', 'guardian cell', 'father cell', 'father phone', 'mother cell', 'mother phone', 'cellphone', 'cell number', 'cell#', 'cell #', 'phone number', 'mobile', 'telephone', 'contact number', 'parent_phone', 'parent 1 phone', 'parent 2 phone'],
    ['jan', 'january'], ['feb', 'february'], ['mar', 'march'], ['apr', 'april'], ['jun', 'june'],
    ['jul', 'july'], ['aug', 'august'], ['sep', 'sept', 'september'], ['oct', 'october'], ['nov', 'november'], ['dec', 'december'],
]


def builtin_aliases_for(field_name: str) -> list[str]:
    nf = norm(field_name)
    out = [field_name]
    for grp in BUILTIN_SYNONYMS:
        if nf in grp:
            out.extend(grp)
    return list(dict.fromkeys(out))


def norm(value):
    return re.sub(r'\s+', ' ', str(value or '').strip()).casefold()


def public_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else str(value)
    if isinstance(value, float):
        return str(value)
    return value


def _coerce_xls_text_cell(raw_val: str):
    s = raw_val.strip()
    if not s:
        return None
    if re.fullmatch(r'0\d+', s) or re.fullmatch(r'\d{11,}', s):
        return s
    if re.fullmatch(r'[+-]?\d+', s):
        try:
            return int(s)
        except ValueError:
            return s
    if re.fullmatch(r'[+-]?\d+\.\d+', s):
        try:
            f = float(s)
            return int(f) if f.is_integer() else f
        except ValueError:
            return s
    return s


def xls_rows(path: Path) -> dict[str, list[list[object]]]:
    raw = path.read_bytes()
    if not raw:
        raise ValueError('Empty .xls file; platform supports .xlsx, .xls and UTF-8 .csv workbooks.')
    # 1. ZIP OpenXML (.xlsx renamed to .xls)
    if raw[:4] == b'PK\x03\x04':
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        result = {}
        try:
            for ws in wb:
                result[ws.title] = [list(row) for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        if result:
            return result
    # 2. Legacy Binary Excel 97-2003 BIFF8 / BIFF5 (.xls)
    if raw[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' or raw[:2] == b'\x09\x08':
        try:
            import xlrd
            book = xlrd.open_workbook(file_contents=raw)
        except Exception as exc:
            raise ValueError('Legacy .xls workbook could not be read; platform supports unencrypted .xlsx, .xls and UTF-8 .csv.') from exc
        result: dict[str, list[list[object]]] = {}
        total = 0
        for sh in book.sheets():
            if sh.nrows > MAX_ROWS or sh.ncols > MAX_COLS:
                raise ValueError(f'Sheet {sh.name} exceeds the row or column limit.')
            total += sh.nrows * sh.ncols
            if total > MAX_CELLS:
                raise ValueError('Workbook exceeds the one-million-cell inspection limit.')
            sheet_rows: list[list[object]] = []
            for r_idx in range(sh.nrows):
                row_vals: list[object] = []
                for c_idx in range(sh.ncols):
                    cell = sh.cell(r_idx, c_idx)
                    ctype = cell.ctype
                    val = cell.value
                    if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row_vals.append(None)
                    elif ctype == xlrd.XL_CELL_DATE:
                        try:
                            dt = xlrd.xldate_as_datetime(val, book.datemode)
                            row_vals.append(dt.strftime('%Y-%m-%d'))
                        except Exception:
                            row_vals.append(str(val))
                    elif ctype == xlrd.XL_CELL_BOOLEAN:
                        row_vals.append(bool(val))
                    elif ctype == xlrd.XL_CELL_NUMBER:
                        fval = float(val)
                        row_vals.append(int(fval) if fval.is_integer() else fval)
                    else:
                        sval = str(val).strip()
                        row_vals.append(sval if sval != '' else None)
                sheet_rows.append(row_vals)
            result[sh.name or f'Sheet{len(result)+1}'] = sheet_rows
        if result:
            return result
    # 3. Text-based .xls exports (HTML table .xls, SpreadsheetML XML 2003 .xls, or TSV/CSV .xls)
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            text = raw.decode('latin-1')
        except Exception as exc:
            raise ValueError('First version supports .xlsx, .xls and UTF-8 .csv sources; unrecognized binary format.') from exc
    lower_head = text[:4096].lower()
    if '<table' in lower_head or '<tr' in lower_head:
        from html.parser import HTMLParser

        class _TableParser(HTMLParser):
            def __init__(self):
                super().__init__()
                self.rows: list[list[object]] = []
                self._cur_row: list[object] | None = None
                self._cur_cell: list[str] | None = None

            def handle_starttag(self, tag, attrs):
                t = tag.lower()
                if t == 'tr':
                    self._cur_row = []
                elif t in ('td', 'th') and self._cur_row is not None:
                    self._cur_cell = []

            def handle_data(self, data):
                if self._cur_cell is not None:
                    self._cur_cell.append(data)

            def handle_endtag(self, tag):
                t = tag.lower()
                if t in ('td', 'th') and self._cur_row is not None and self._cur_cell is not None:
                    self._cur_row.append(_coerce_xls_text_cell(''.join(self._cur_cell)))
                    self._cur_cell = None
                elif t == 'tr' and self._cur_row is not None:
                    if any(v is not None for v in self._cur_row):
                        self.rows.append(self._cur_row)
                    self._cur_row = None

        parser = _TableParser()
        parser.feed(text)
        if parser.rows:
            return {'Sheet1': parser.rows}
    if '<workbook' in lower_head or '<worksheet' in lower_head:
        import xml.etree.ElementTree as ET
        try:
            root = ET.fromstring(text)
            sheets: dict[str, list[list[object]]] = {}
            for ws_el in root.iter():
                if ws_el.tag.endswith('Worksheet'):
                    sh_name = next((v for k, v in ws_el.attrib.items() if k.endswith('Name')), f'Sheet{len(sheets)+1}')
                    sh_rows: list[list[object]] = []
                    for r_el in ws_el.iter():
                        if r_el.tag.endswith('Row'):
                            r_vals: list[object] = []
                            for c_el in r_el:
                                if c_el.tag.endswith('Cell'):
                                    idx_attr = next((v for k, v in c_el.attrib.items() if k.endswith('Index')), None)
                                    if idx_attr and str(idx_attr).isdigit():
                                        target_idx = int(idx_attr) - 1
                                        while len(r_vals) < target_idx:
                                            r_vals.append(None)
                                    cell_text = ''.join(c_el.itertext())
                                    r_vals.append(_coerce_xls_text_cell(cell_text))
                            sh_rows.append(r_vals)
                    if sh_rows:
                        sheets[sh_name] = sh_rows
            if sheets:
                return sheets
        except Exception:
            pass
    # 4. Tab/comma/semicolon delimited table saved with .xls extension (must have >=2 lines and a delimiter)
    non_empty_lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(non_empty_lines) >= 2 and any(d in non_empty_lines[0] for d in ('\t', ',', ';')):
        return {'Sheet1': csv_rows(path)}
    raise ValueError('First version supports .xlsx, .xls and UTF-8 .csv sources; templates must be .xlsx. Legacy, macro-enabled and encrypted workbooks are not supported.')


def check_file(path: Path):
    if path.stat().st_size > MAX_BYTES:
        raise ValueError('File exceeds the first-version 20 MiB limit.')
    if path.suffix.lower() not in {'.xlsx', '.xls', '.csv'}:
        raise ValueError('First version supports .xlsx, .xls and UTF-8 .csv sources; templates must be .xlsx. Legacy, macro-enabled and encrypted workbooks are not supported.')
    if path.suffix.lower() == '.xls':
        xls_rows(path)
        return
    if path.suffix.lower() == '.xlsx':
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(x.file_size for x in entries) > 80 * 1024 * 1024:
                    raise ValueError('Workbook expands beyond the safe inspection limit.')
                if any('vbaProject' in x.filename for x in entries):
                    raise ValueError('Macro content is not supported or executed.')
                if any(x.filename.startswith('/') or '..' in Path(x.filename).parts for x in entries):
                    raise ValueError('Workbook contains unsafe archive paths.')
        except zipfile.BadZipFile as exc:
            raise ValueError('This is not a readable, unencrypted .xlsx workbook.') from exc


def csv_rows(path):
    try:
        text = path.read_text(encoding='utf-8-sig')
    except UnicodeDecodeError as exc:
        raise ValueError('CSV sources must be UTF-8; convert this source explicitly before import.') from exc
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        if len(row) > MAX_COLS:
            raise ValueError('Source exceeds the 200-column limit.')
        rows.append([cell if cell != '' else None for cell in row])
        if len(rows) > MAX_ROWS:
            raise ValueError('Source exceeds the 20,000-row per-sheet limit.')
    return rows


_MONTH_MAP = {
    'jan': 1, 'january': 1, 'feb': 2, 'february': 2, 'mar': 3, 'march': 3,
    'apr': 4, 'april': 4, 'may': 5, 'jun': 6, 'june': 6,
    'jul': 7, 'july': 7, 'aug': 8, 'august': 8, 'sep': 9, 'sept': 9, 'september': 9,
    'oct': 10, 'october': 10, 'nov': 11, 'november': 11, 'dec': 12, 'december': 12,
}


def _extract_dob_from_sa_id_str(raw_val) -> tuple[str | None, str | None]:
    if raw_val is None or isinstance(raw_val, bool):
        return None, None
    s = re.sub(r'\.0+$', '', str(raw_val).strip())
    if re.search(r'[\r\n]', s):
        first_line = next((ln.strip() for ln in re.split(r'[\r\n]+', s) if re.search(r'\d{5,13}', ln)), None)
        if first_line:
            s = first_line
    if re.fullmatch(r'\d+(?:\.\d+)?[eE]\+\d+', s):
        try:
            s = str(int(round(float(s))))
        except Exception:
            pass
    s_clean = re.sub(r'[\s-]', '', s)
    if len(s_clean) >= 20 and s_clean.isdigit():
        m13 = re.search(r'\b(\d{13})\b', s)
        s_clean = m13.group(1) if m13 else s_clean[:13]
    if not s_clean or not re.fullmatch(r'\d+', s_clean):
        return None, None
    digits = s_clean
    # Support 6-digit YYMMDD (e.g. 121007 -> 2012-10-07) and 5-digit zero-stripped YMMDD for 2005-2009 (e.g. 90524 -> 090524 -> 2009-05-24)
    if len(digits) == 6 or (len(digits) == 5 and digits[0] in '56789'):
        d6 = digits.zfill(6)
        yy, mm, dd = int(d6[:2]), int(d6[2:4]), int(d6[4:6])
        if 1 <= mm <= 12 and 1 <= dd <= 31:
            century = 2000 if yy <= 26 else 1900
            return f'{century + yy:04d}-{mm:02d}-{dd:02d}', None
        return None, None
    if 10 <= len(digits) <= 12:
        padded = digits.zfill(13)
        mm_p, dd_p = int(padded[2:4]), int(padded[4:6])
        if 1 <= mm_p <= 12 and 1 <= dd_p <= 31:
            digits = padded
    if len(digits) != 13:
        return None, None
    yy, mm, dd = int(digits[:2]), int(digits[2:4]), int(digits[4:6])
    if not (1 <= mm <= 12 and 1 <= dd <= 31):
        return None, None
    century = 2000 if yy <= 26 else 1900
    return f'{century + yy:04d}-{mm:02d}-{dd:02d}', digits


def _normalize_date_str(raw_val) -> str | None:
    if raw_val is None or isinstance(raw_val, bool):
        return None
    if isinstance(raw_val, (datetime, date)):
        dt_val = raw_val.date() if isinstance(raw_val, datetime) else raw_val
        if 1900 <= dt_val.year <= 2030:
            return dt_val.strftime('%Y-%m-%d')
        # Recover raw 6-digit YYMMDD integer (e.g. 121007 -> year 2231) when a Date-formatted cell converted YYMMDD into a date!
        serial = (dt_val - date(1899, 12, 30)).days
        dob, _ = _extract_dob_from_sa_id_str(serial)
        return dob
    if isinstance(raw_val, (int, float, Decimal)):
        num = float(raw_val)
        if 20000 <= num <= 50000:
            from datetime import timedelta
            dt = date(1899, 12, 30) + timedelta(days=int(round(num)))
            if 1955 <= dt.year <= 2030:
                return dt.strftime('%Y-%m-%d')
        elif 50101 <= num <= 991231 and int(round(num)) == num:
            dob, _ = _extract_dob_from_sa_id_str(int(round(num)))
            if dob:
                return dob
    s = str(raw_val).strip()
    if not s:
        return None
    m_iso = re.match(r'^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[T\s].*)?$', s)
    if m_iso:
        y, m, d = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
        if 1 <= m <= 12 and 1 <= d <= 31:
            if 1900 <= y <= 2035:
                return f'{y:04d}-{m:02d}-{d:02d}'
            if y > 2035:
                try:
                    serial = (date(y, m, d) - date(1899, 12, 30)).days
                    dob, _ = _extract_dob_from_sa_id_str(serial)
                    if dob:
                        return dob
                except Exception:
                    pass
    m_dmy4 = re.match(r'^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})(?:\s.*)?$', s)
    if m_dmy4:
        a, b, y = int(m_dmy4.group(1)), int(m_dmy4.group(2)), int(m_dmy4.group(3))
        if 1900 <= y <= 2035:
            if 1 <= b <= 12 and 1 <= a <= 31:
                return f'{y:04d}-{b:02d}-{a:02d}'
            if 1 <= a <= 12 and 1 <= b <= 31:
                return f'{y:04d}-{a:02d}-{b:02d}'
    m_dmy2 = re.match(r'^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2})$', s)
    if m_dmy2:
        a, b, c = int(m_dmy2.group(1)), int(m_dmy2.group(2)), int(m_dmy2.group(3))
        if 1 <= b <= 12 and 1 <= a <= 31:
            yr = 2000 + c if c <= 26 else 1900 + c
            return f'{yr:04d}-{b:02d}-{a:02d}'
        if 1 <= b <= 12 and 1 <= c <= 31:
            yr = 2000 + a if a <= 26 else 1900 + a
            return f'{yr:04d}-{b:02d}-{c:02d}'
    m_mon1 = re.match(r'^(\d{1,2})[\s./-]+([A-Za-z]{3,9})[\s,./-]+(\d{2,4})$', s)
    if m_mon1:
        d = int(m_mon1.group(1))
        mon = _MONTH_MAP.get(m_mon1.group(2).lower())
        yr = int(m_mon1.group(3))
        if yr < 100:
            yr = 2000 + yr if yr <= 26 else 1900 + yr
        if mon and 1 <= d <= 31 and 1900 <= yr <= 2035:
            return f'{yr:04d}-{mon:02d}-{d:02d}'
    m_mon2 = re.match(r'^([A-Za-z]{3,9})[\s./-]+(\d{1,2})[\s,./-]+(\d{2,4})$', s)
    if m_mon2:
        mon = _MONTH_MAP.get(m_mon2.group(1).lower())
        d = int(m_mon2.group(2))
        yr = int(m_mon2.group(3))
        if yr < 100:
            yr = 2000 + yr if yr <= 26 else 1900 + yr
        if mon and 1 <= d <= 31 and 1900 <= yr <= 2035:
            return f'{yr:04d}-{mon:02d}-{d:02d}'
    if re.fullmatch(r'(?:19|20)\d{6}', s):
        y, m, d = int(s[:4]), int(s[4:6]), int(s[6:8])
        if 1 <= m <= 12 and 1 <= d <= 31:
            return f'{y:04d}-{m:02d}-{d:02d}'
    if re.fullmatch(r'\d{5}(?:\.0+)?', s):
        serial = int(round(float(s)))
        if 20000 <= serial <= 50000:
            from datetime import timedelta
            dt = date(1899, 12, 30) + timedelta(days=serial)
            if 1955 <= dt.year <= 2030:
                return dt.strftime('%Y-%m-%d')
    dob, _ = _extract_dob_from_sa_id_str(s)
    return dob


def _title_case_person(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    if s.isupper() or s.islower():
        return ' '.join(w[:1].upper() + w[1:].lower() for w in s.split() if w)
    return s


def _split_surname_and_first_names(raw_val) -> tuple[str | None, str | None]:
    if raw_val is None:
        return None, None
    s = re.sub(r'\s+', ' ', str(raw_val).replace('\r', ' ').replace('\n', ' ')).strip()
    if not s or len(s) < 3 or re.search(r'\d', s):
        return None, None
    if ',' in s:
        parts = [p.strip() for p in s.split(',') if p.strip()]
        if len(parts) >= 2:
            return _title_case_person(parts[0]), _title_case_person(' '.join(parts[1:]))
    parts = [p for p in s.split(' ') if p]
    if len(parts) >= 2:
        if len(parts) >= 4 and parts[0].lower() in {'van', 'de'} and parts[1].lower() in {'der', 'den', 'la'}:
            return _title_case_person(' '.join(parts[:3])), _title_case_person(' '.join(parts[3:]))
        if len(parts) >= 3 and parts[0].lower() in {'van', 'de', 'du', 'le', 'la', 'mc', 'mac', 'st'}:
            return _title_case_person(' '.join(parts[:2])), _title_case_person(' '.join(parts[2:]))
        return _title_case_person(parts[0]), _title_case_person(' '.join(parts[1:]))
    return None, None


_SECTION_HEADER_CELL_RE = re.compile(
    r'^((?:learner\s*|student\s*|parent\s*|guardian\s*|father\s*|mother\s*)?'
    r'(?:number|no\.?|#|id|student\s*id|learner\s*id|accession(?:\s*number|\s*no\.?)?|'
    r'admission(?:\s*number|\s*no\.?)?|sa\s*id(?:\s*number)?|id\s*number|identity\s*number|'
    r'lurits(?:\s*no\.?)?|cemis(?:\s*no\.?)?|surname|last\s*name|lastname|first\s*name|firstname|'
    r'name|names|given\s*name|forename|full\s*name|learners?\s+in\s+alphabetical\s+order.*|'
    r'.*surname\s+(?:and|&)\s+first\s*names?.*|gender|sex|male\s*\/\s*female|grade|gr\.?|class|birth\s*date|'
    r'date\s*of\s*birth|dob|d\.o\.b\.?|birthdate|age|number\s+of\s+years.*|number\s+of\s+days.*|'
    r'cell\s*#?|cell\s*no\.?|cellphone|'
    r'phone(?:\s*number)?|tel(?:ephone)?|mobile|contact(?:\s*number)?|email|e-mail|email\s*address|'
    r'home\s*language|first\s*additional\s*language|second\s*additional\s*language|subjects?|'
    r'subject\s*list|enrolled\s*subjects|learner\s*average.*|learner\s*rank|learner\s*total|code|codes\s*:.*|'
    r'deviation\s*comment|.*report\s*general\s*comment.*|.*\(\s*gr(?:ade)?\.?\s*\d{1,2}\s*\)|'
    r'accounting|business\s*studies|economics|mathematics|mathematical\s*literacy|physical\s*sciences|'
    r'life\s*sciences|life\s*orientation|geography|history|tourism|consumer\s*studies|'
    r'computer\s*applications\s*technology|natural\s*sciences|social\s*sciences|technology|'
    r'creative\s*arts|economic\s*(?:and\s*)?management.*|ems))$',
    re.I,
)


def _is_banner_or_metadata_row(row: list[object]) -> bool:
    filled = [v for v in (row or []) if v is not None and str(v).strip() != '']
    if not filled:
        return True
    if len(filled) <= 3:
        joined = ' | '.join(re.sub(r'\s+', ' ', str(v).strip()) for v in filled)
        if re.search(
            r'\b(department\s+of\s+education|internal\s+school|mark\s+schedule|school\s*:|'
            r'district\s*:|circuit\s*:|emis\s*:|for\s*:\s*grade|report\s+generated|prepared\s+by|'
            r'page\s+\d+\s+of|summary\s+report|total\s+learners|learner\s+information|parent\s+information)\b',
            joined,
            re.I,
        ):
            return True
    return False


def _is_mid_sheet_school_header_row(row: list[object], prev_row: list[object] | None = None) -> bool:
    max_len = max(len(row or []), len(prev_row or []))
    if max_len < 3:
        return False
    row_filled = [v for v in (row or []) if v is not None and str(v).strip() != '']
    if len(row_filled) < 2:
        return False
    first_cell = row[0] if row else None
    if isinstance(first_cell, (int, float)) or (first_cell is not None and re.fullmatch(r'\d+', str(first_cell).strip())):
        return False
    composite: list[object] = []
    for c_i in range(max_len):
        cur = row[c_i] if row and c_i < len(row) else None
        cur_s = str(cur).strip() if cur is not None else ''
        prv = prev_row[c_i] if prev_row and c_i < len(prev_row) else None
        prv_s = str(prv).strip() if prv is not None else ''
        if cur_s and prv_s and len(prv_s) < 40 and not _SECTION_HEADER_CELL_RE.match(cur_s) and _SECTION_HEADER_CELL_RE.match(f'{prv_s} {cur_s}'):
            composite.append(f'{prv_s} {cur_s}')
        elif cur_s:
            composite.append(cur)
        elif prv_s and len(prv_s) < 45 and _SECTION_HEADER_CELL_RE.match(prv_s):
            composite.append(prv)
    if len(composite) < 3:
        return False
    hits = 0
    non_num_filled = 0
    for v in composite:
        if isinstance(v, (int, float, Decimal)):
            continue
        s = re.sub(r'\s+', ' ', str(v).strip())
        if not s or re.fullmatch(r'\d+', s):
            continue
        non_num_filled += 1
        if '@' in s or re.fullmatch(r'(?:\+27|0)[1-9]\d{7,8}', re.sub(r'[\s()-]', '', s)):
            return False
        if re.fullmatch(r'\d{10,13}', re.sub(r'\s+', '', s)) or re.fullmatch(r'\d+[.,]\d+', s):
            return False
        if re.fullmatch(r'(?:grade|gr\.?)\s*\d{1,2}[a-z]?', s, re.I) or re.match(r'^[A-Z][A-Z\s\'-]+,\s*[A-Za-z]+', s):
            return False
        if _SECTION_HEADER_CELL_RE.match(s):
            hits += 1
    return hits >= 3 and non_num_filled >= 2 and hits >= non_num_filled * 0.45


def _split_and_normalize_sheet_sections(sheet_name: str, rows: list[list[object]]) -> dict[str, list[list[object]]]:
    if not rows or len(rows) < 2 or norm(sheet_name) in RULE_SHEETS:
        return {sheet_name: rows}
    first_h_idx = 0
    for idx, r in enumerate(rows[:35]):
        prev_r = rows[idx - 1] if idx > 0 else None
        if _is_mid_sheet_school_header_row(r, prev_r):
            first_h_idx = idx
            break
    header_indices = [first_h_idx]
    for r_idx in range(first_h_idx + 2, len(rows)):
        if _is_mid_sheet_school_header_row(rows[r_idx], rows[r_idx - 1]):
            if r_idx - header_indices[-1] <= 3:
                header_indices[-1] = r_idx
                continue
            header_indices.append(r_idx)
    if len(header_indices) == 1:
        return {sheet_name: _normalize_school_sheet_rows(rows, sheet_name=sheet_name)}
    out_sections: dict[str, list[list[object]]] = {}
    for s_i, h_idx in enumerate(header_indices):
        section_banner_grade: str | None = None
        scan_min = 0 if s_i == 0 else max(header_indices[s_i - 1] + 1, h_idx - 18)
        for b_idx in range(h_idx, scan_min - 1, -1):
            for cell in (rows[b_idx] or []):
                if not isinstance(cell, str):
                    continue
                gm = re.search(r'(?:for\s*:\s*grade\s*|grade\s*:\s*|\(\s*gr(?:ade)?\.?\s*|\bgrade\s+|gr\.?\s*)0*(8|9|10|11|12)\b', cell, re.I)
                if gm:
                    section_banner_grade = f'Grade {int(gm.group(1))}'
                    break
            if section_banner_grade:
                break
        start_idx = 0 if s_i == 0 else max(header_indices[s_i - 1] + 2, h_idx - 6)
        if s_i > 0:
            back = h_idx - 1
            min_back = max(header_indices[s_i - 1] + 2, h_idx - 14)
            while back >= min_back:
                r_back = rows[back] or []
                is_data = any(
                    isinstance(c, str) and (
                        re.match(r'^[A-Z][A-Z\s\'-]+,\s*[A-Za-z]+', c.strip())
                        or re.fullmatch(r'\d{10,13}', re.sub(r'\s+', '', c.strip()))
                    )
                    for c in r_back
                )
                if is_data:
                    break
                start_idx = back
                back -= 1
        end_idx = len(rows)
        if s_i + 1 < len(header_indices):
            next_start = header_indices[s_i + 1]
            back = next_start - 1
            min_back = max(h_idx + 2, next_start - 12)
            while back >= min_back and _is_banner_or_metadata_row(rows[back]):
                next_start = back
                back -= 1
            end_idx = next_start
        sub_rows = [r for r in rows[start_idx:end_idx] if any(v is not None and str(v).strip() != '' for v in (r or []))]
        if not sub_rows:
            continue
        sec_key = sheet_name if s_i == 0 else f'{sheet_name} (Section {s_i + 1})'
        out_sections[sec_key] = _normalize_school_sheet_rows(sub_rows, sheet_name=sheet_name, section_banner_grade=section_banner_grade)
    return out_sections or {sheet_name: _normalize_school_sheet_rows(rows, sheet_name=sheet_name)}


def _normalize_school_sheet_rows(
    rows: list[list[object]],
    sheet_name: str | None = None,
    section_banner_grade: str | None = None,
) -> list[list[object]]:
    if not rows or len(rows) < 2:
        return rows
    # Scan sheet tab name, section banner grade, and top 20 rows for Grade 8..12 banner
    banner_grade: str | None = section_banner_grade
    if not banner_grade and sheet_name:
        sm = re.search(r'\b(?:grade|gr\.?)\s*0*(8|9|10|11|12)\b', sheet_name, re.I)
        if sm:
            banner_grade = f'Grade {int(sm.group(1))}'
    for r in rows[:20]:
        if banner_grade:
            break
        row_joined = ' '.join(str(cell).strip() for cell in (r or []) if cell is not None and str(cell).strip())
        m = re.search(r'(?:for\s*:\s*|grade\s*:\s*)?(?:grade|gr\.?)\s*0*(\d{1,2})\b', row_joined, re.I) or re.search(r'\(\s*gr(?:ade)?\.?\s*0*(\d{1,2})\s*\)', row_joined, re.I)
        if m:
            banner_grade = f'Grade {int(m.group(1))}'
            break

    header_idx: int | None = None
    mode: str | None = None
    for idx, r in enumerate(rows[:35]):
        max_c = max([len(rows[k] or []) for k in range(max(0, idx - 4), idx + 1)] or [0])
        comp_norms: list[str] = []
        for c_i in range(max_c):
            for k in range(idx, max(-1, idx - 5), -1):
                if k < len(rows) and c_i < len(rows[k] or []):
                    v = rows[k][c_i]
                    if v is not None and str(v).strip():
                        comp_norms.append(norm(v))
                        break
        norms = [norm(v) for v in (r or []) if v is not None and str(v).strip()]
        if not norms:
            continue
        has_learner_name_col = any(
            'learner name' in n
            or 'student name' in n
            or 'learners in alphabetical order' in n
            or 'surname and first name' in n
            for n in comp_norms
        )
        if has_learner_name_col and (
            any('accession' in n or 'admission' in n or n in {'date of birth', 'dob', 'birth date', 'birthdate', 'no', 'no.', 'number', 'gender', 'male female'} for n in comp_norms)
        ):
            header_idx = idx
            mode = 'gde_marks'
            break
        name_cnt = sum(1 for n in norms if n in {'first name', 'firstname', 'name', 'names', 'given name'})
        sur_cnt = sum(1 for n in norms if n in {'last name', 'lastname', 'surname'})
        if name_cnt >= 2 and sur_cnt >= 2:
            header_idx = idx
            mode = 'bio_parents'
            break

    if header_idx is None or mode is None:
        return rows

    out_rows = [list(r) for r in rows]
    # Check if header_idx + 1 is a sub-header row (e.g. when header_idx has Number/Admission Number/Learner's Progress and header_idx + 1 has the 9 subjects)
    if mode == 'gde_marks' and header_idx + 1 < len(out_rows):
        nxt_r = out_rows[header_idx + 1] or []
        nxt_first = nxt_r[0] if nxt_r else None
        if not isinstance(nxt_first, (int, float)) and not (nxt_first is not None and re.fullmatch(r'\d+', str(nxt_first).strip())):
            nxt_hits = sum(1 for v in nxt_r if v is not None and str(v).strip() and _SECTION_HEADER_CELL_RE.match(re.sub(r'\s+', ' ', str(v).strip())))
            if nxt_hits >= 3:
                header_idx = header_idx + 1
    h_row = list(out_rows[header_idx])
    while h_row and (h_row[-1] is None or str(h_row[-1]).strip() == ''):
        h_row.pop()

    if mode == 'gde_marks':
        # Clean long comment header and composite any blank header cells from up to 5 rows above header_idx
        clean_h: list[str] = []
        max_h_cols = max([len(out_rows[k] or []) for k in range(max(0, header_idx - 3), min(len(out_rows), header_idx + 5))] or [len(h_row)])
        while len(h_row) < max_h_cols:
            h_row.append(None)
        while h_row and all(
            (k >= len(out_rows) or len(out_rows[k]) < len(h_row) or out_rows[k][len(h_row) - 1] is None or str(out_rows[k][len(h_row) - 1]).strip() == '')
            for k in range(max(0, header_idx - 2), min(len(out_rows), header_idx + 8))
        ):
            h_row.pop()
        for c_i, cell in enumerate(h_row):
            hs = re.sub(r'\s+', ' ', str(cell or '').strip())
            if not hs or re.match(r'^learner\'?s\s+progress$', hs, re.I):
                for k in range(header_idx - 1, max(-1, header_idx - 6), -1):
                    if c_i < len(out_rows[k]):
                        pk = re.sub(r'\s+', ' ', str(out_rows[k][c_i] or '').strip())
                        if pk and len(pk) <= 140 and not re.match(r'^learner\'?s\s+progress$', pk, re.I) and not re.search(r'\b(department\s+of\s+education|mark\s+schedule|school\s*:|district\s*:|for\s*:\s*grade|grade\s*:\s*grade)\b', pk, re.I):
                            hs = pk
                            break
            if re.search(r'report\s+general\s+comment', hs, re.I):
                hs = 'Report General Comment'
            elif re.search(r'codes\s*:.*deviation\s+comment', hs, re.I):
                hs = 'Deviation Comment'
            elif re.search(r'learners?\s+in\s+alphabetical\s+order|surname\s+(?:and|&)\s+first\s*names?', hs, re.I):
                hs = 'Learner Name'
            elif re.match(r'^(male\s*\/\s*female|m\s*\/\s*f)$', hs, re.I):
                hs = 'Gender'
            elif re.match(r'^economic\s+management$', hs, re.I):
                hs = 'Economic and Management Sciences (EMS)'
            if not hs:
                hs = f'Col{c_i + 1}'
            if not banner_grade:
                gm = re.search(r'\(\s*gr(?:ade)?\.?\s*0*(\d{1,2})\s*\)', hs, re.I)
                if gm:
                    banner_grade = f'Grade {int(gm.group(1))}'
            clean_h.append(hs)
        acc_idx = next((i for i, h in enumerate(clean_h) if 'accession' in norm(h) or 'admission' in norm(h)), None)
        dob_col_idx = next((i for i, h in enumerate(clean_h) if norm(h) in {'date of birth', 'dob', 'birth date', 'birthdate'}), None)
        lname_idx = next((i for i, h in enumerate(clean_h) if norm(h) in {'learner name', 'student name', 'full name'}), None)
        has_grade = any(norm(h) == 'grade' for h in clean_h)
        clean_h.extend(['Last Name', 'First Name'])
        if dob_col_idx is None:
            clean_h.append('Date of Birth')
        if not has_grade:
            clean_h.append('Grade')
        out_rows[header_idx] = clean_h
        current_grade = banner_grade
        for r_i in range(header_idx + 1, len(out_rows)):
            r = list(out_rows[r_i])
            if not any(v is not None and str(v).strip() != '' for v in r):
                continue
            for cell_v in r[:6]:
                if isinstance(cell_v, str):
                    gm_row = re.search(r'(?:for\s*:\s*grade\s*|grade\s*:\s*(?:grade\s*)?|\(\s*gr(?:ade)?\.?\s*)0*(8|9|10|11|12)\b', cell_v, re.I)
                    if gm_row:
                        current_grade = f'Grade {int(gm_row.group(1))}'
            while len(r) < len(h_row):
                r.append(None)
            r = r[:len(h_row)]
            # Split "SURNAME, FirstName, MiddleName" OR space-separated "SURNAME FIRSTNAMES"
            ln_val, fn_val = None, None
            if lname_idx is not None and lname_idx < len(r) and r[lname_idx] is not None:
                ln_val, fn_val = _split_surname_and_first_names(r[lname_idx])
            dob_val = None
            if dob_col_idx is not None and dob_col_idx < len(r) and r[dob_col_idx] is not None:
                norm_d = _normalize_date_str(r[dob_col_idx])
                if norm_d:
                    r[dob_col_idx] = norm_d
                    dob_val = norm_d
            if acc_idx is not None and acc_idx < len(r) and r[acc_idx] is not None:
                ext_dob, padded_id = _extract_dob_from_sa_id_str(r[acc_idx])
                if ext_dob and not dob_val:
                    dob_val = ext_dob
                    if dob_col_idx is not None and dob_col_idx < len(r):
                        r[dob_col_idx] = ext_dob
                if padded_id:
                    r[acc_idx] = padded_id
            if not dob_val:
                for sc_i in range(min(5, len(r))):
                    if sc_i == lname_idx:
                        continue
                    ext_dob, _ = _extract_dob_from_sa_id_str(r[sc_i])
                    if not ext_dob:
                        ext_dob = _normalize_date_str(r[sc_i])
                    if ext_dob and 1990 <= int(ext_dob[:4]) <= 2022:
                        dob_val = ext_dob
                        if dob_col_idx is not None and dob_col_idx < len(r):
                            r[dob_col_idx] = ext_dob
                        break
            r.extend([ln_val, fn_val])
            if dob_col_idx is None:
                r.append(dob_val)
            if not has_grade:
                r.append(current_grade)
            out_rows[r_i] = r
        return out_rows

    if mode == 'bio_parents':
        first_roles = {'first name', 'firstname', 'name', 'names', 'given name', 'forename'}
        last_roles = {'last name', 'lastname', 'surname', 'family name'}
        phone_roles = {'phone', 'cell', 'cell#', 'cell #', 'cell no', 'cellphone', 'mobile', 'tel', 'telephone', 'contact', 'contact number'}
        email_roles = {'email', 'e-mail', 'email address', 'mail', 'contact email'}
        def _is_dob_col(n_str: str) -> bool:
            if not n_str or 'place' in n_str or 'country' in n_str or 'certificate' in n_str:
                return False
            return n_str in {'date of birth', 'dob', 'birth date', 'birthdate', 'date_of_birth', 'learner dob', 'd.o.b', 'd.o.b.', 'birth'} or bool(re.search(r'(?:^|[\s_.-])(dob|birth\s*date|birthdate|date\s*of\s*birth)(?:$|[\s_.-])', n_str))
        first_parent_col = 999
        fc, lc = 0, 0
        for i, cell in enumerate(h_row):
            n = norm(cell)
            if n in first_roles:
                fc += 1
                if fc == 2 and i < first_parent_col:
                    first_parent_col = i
            if n in last_roles:
                lc += 1
                if lc == 2 and i < first_parent_col:
                    first_parent_col = i
        email_cols_after_parent = sum(1 for i, cell in enumerate(h_row) if i >= first_parent_col and (norm(cell) in email_roles or bool(re.search(r'\b(parent|guardian|father|mother)\b.*\b(email|e-mail)\b', norm(cell)))))
        f_cnt, l_cnt, p_cnt, e_cnt, dob_cnt = 0, 0, 0, 0, 0
        raw_clean_h = []
        for i, cell in enumerate(h_row):
            raw_s = re.sub(r'\s+', ' ', str(cell or '').strip()) or f'Col{i + 1}'
            n = norm(raw_s)
            if n in first_roles:
                f_cnt += 1
                raw_clean_h.append('First Name' if f_cnt == 1 else f'Parent {f_cnt - 1} First Name')
            elif n in last_roles:
                l_cnt += 1
                raw_clean_h.append('Last Name' if l_cnt == 1 else f'Parent {l_cnt - 1} Last Name')
            elif n in phone_roles:
                if i < first_parent_col:
                    raw_clean_h.append('Learner Phone')
                else:
                    p_cnt += 1
                    raw_clean_h.append(f'Parent {p_cnt} Phone')
            elif n in email_roles or bool(re.search(r'\b(parent|guardian|father|mother)\b.*\b(email|e-mail)\b', n)):
                is_explicit_parent_em = bool(re.search(r'\b(parent|guardian|father|mother)\b', n))
                if not is_explicit_parent_em and (i < first_parent_col or email_cols_after_parent == 1):
                    raw_clean_h.append('Learner Email')
                else:
                    e_cnt += 1
                    raw_clean_h.append(f'Parent {e_cnt} Email')
            elif _is_dob_col(n):
                is_explicit_parent_dob = bool(re.search(r'\b(parent|guardian|father|mother)\b', n))
                if not is_explicit_parent_dob and dob_cnt == 0:
                    dob_cnt += 1
                    raw_clean_h.append('Date of Birth')
                else:
                    dob_cnt += 1
                    raw_clean_h.append(f'Parent {max(1, dob_cnt - 1)} Date of Birth')
            else:
                raw_clean_h.append(raw_s)
        seen_h: dict[str, int] = {}
        clean_h: list[str] = []
        for h_item in raw_clean_h:
            cnt = seen_h.get(h_item, 0) + 1
            seen_h[h_item] = cnt
            clean_h.append(h_item if cnt == 1 else f'{h_item} ({cnt})')
        idx_map = {h: i for i, h in enumerate(clean_h)}
        has_explicit_parent_emails = ('Parent 1 Email' in idx_map) or ('Parent 2 Email' in idx_map)
        has_dob_col = 'Date of Birth' in idx_map
        sa_id_idxs = [i for i, h in enumerate(clean_h) if bool(re.search(r'\b(id|accession|identity|lurits|cemis)\b', h, re.I)) and not bool(re.search(r'parent|guardian|father|mother', h, re.I))]
        clean_h.extend(['Parent First Name', 'Parent Last Name', 'Parent Phone'])
        if has_explicit_parent_emails:
            clean_h.append('Parent Email')
        if not has_dob_col:
            clean_h.append('Date of Birth')
        out_rows[header_idx] = clean_h
        for r_i in range(header_idx + 1, len(out_rows)):
            r = list(out_rows[r_i])
            if not any(v is not None and str(v).strip() != '' for v in r):
                continue
            while len(r) < len(h_row):
                r.append(None)
            r = r[:len(h_row)]

            def _cell_str(col_name: str) -> str:
                ci = idx_map.get(col_name)
                if ci is None or ci >= len(r) or r[ci] is None:
                    return ''
                s = str(r[ci]).strip()
                if col_name.endswith('Phone') and re.fullmatch(r'0+', s):
                    return ''
                return s

            # If a row in bio_parents has a comma-separated "SURNAME, FirstName" in First Name / Last Name (e.g. a mid-sheet table where Col B is 6-digit YYMMDD and Col C is "MASHABA, Boitumelo, Faith"), split it cleanly!
            fn_ci, ln_ci = idx_map.get('First Name'), idx_map.get('Last Name')
            fn_cur = _cell_str('First Name')
            ln_cur = _cell_str('Last Name')
            if (not fn_cur or not ln_cur or re.search(r'\d', fn_cur) or re.search(r'\d', ln_cur) or ',' in fn_cur or ',' in ln_cur):
                for sc_i in range(min(5, len(r))):
                    sv = str(r[sc_i] or '').strip()
                    if re.match(r'^[A-Za-z][A-Za-z\s\'-]+,\s*[A-Za-z]+', sv):
                        parts = [p.strip() for p in sv.split(',') if p.strip()]
                        if len(parts) >= 2:
                            if ln_ci is not None and ln_ci < len(r):
                                r[ln_ci] = _title_case_person(parts[0])
                            if fn_ci is not None and fn_ci < len(r):
                                r[fn_ci] = _title_case_person(' '.join(parts[1:]))
                            break

            # Normalize learner Date of Birth on the row (from Date of Birth column, SA ID column, or date cell scan)
            resolved_dob = None
            if has_dob_col:
                dob_ci = idx_map['Date of Birth']
                if dob_ci < len(r) and r[dob_ci] is not None:
                    resolved_dob = _normalize_date_str(r[dob_ci])
            if not resolved_dob:
                for sid_i in sa_id_idxs:
                    if sid_i < len(r) and r[sid_i] is not None:
                        resolved_dob, padded_id = _extract_dob_from_sa_id_str(r[sid_i])
                        if padded_id:
                            r[sid_i] = padded_id
                        if resolved_dob:
                            break
            if not resolved_dob:
                for c_idx, cell_v in enumerate(r):
                    if c_idx >= first_parent_col and clean_h[c_idx].startswith(('Parent 1 ', 'Parent 2 ')):
                        continue
                    cand_d = _normalize_date_str(cell_v)
                    if cand_d and 1990 <= int(cand_d[:4]) <= 2022:
                        resolved_dob = cand_d
                        break
            if has_dob_col and resolved_dob:
                r[idx_map['Date of Birth']] = resolved_dob

            p1_fn, p1_ln, p1_ph, p1_em = _cell_str('Parent 1 First Name'), _cell_str('Parent 1 Last Name'), _cell_str('Parent 1 Phone'), _cell_str('Parent 1 Email')
            p2_fn, p2_ln, p2_ph, p2_em = _cell_str('Parent 2 First Name'), _cell_str('Parent 2 Last Name'), _cell_str('Parent 2 Phone'), _cell_str('Parent 2 Email')
            if p1_fn or p1_ln:
                r.extend([_title_case_person(p1_fn), _title_case_person(p1_ln), p1_ph or None])
                if has_explicit_parent_emails:
                    r.append(p1_em or None)
            elif p2_fn or p2_ln:
                r.extend([_title_case_person(p2_fn), _title_case_person(p2_ln), p2_ph or None])
                if has_explicit_parent_emails:
                    r.append(p2_em or None)
            else:
                r.extend([None, None, p1_ph or p2_ph or None])
                if has_explicit_parent_emails:
                    r.append(p1_em or p2_em or None)
            if not has_dob_col:
                r.append(resolved_dob)
            out_rows[r_i] = r
        return out_rows

    return rows


def book_rows(path: Path, data_only=False):
    check_file(path)
    if path.suffix.lower() == '.csv':
        out_csv: dict[str, list[list[object]]] = {}
        for k, v in {'CSV': csv_rows(path)}.items():
            out_csv.update(_split_and_normalize_sheet_sections(k, v))
        return out_csv
    if path.suffix.lower() == '.xls':
        out_xls: dict[str, list[list[object]]] = {}
        for k, v in xls_rows(path).items():
            out_xls.update(_split_and_normalize_sheet_sections(k, v))
        return out_xls
    try:
        wb = load_workbook(path, read_only=True, data_only=data_only, keep_links=False)
    except Exception as exc:
        raise ValueError('Workbook could not be read safely. Check its format and encryption.') from exc
    result = {}
    total = 0
    try:
        for ws in wb:
            if ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS:
                raise ValueError(f'Sheet {ws.title} exceeds the row or column limit.')
            total += ws.max_row * ws.max_column
            if total > MAX_CELLS:
                raise ValueError('Workbook exceeds the one-million-cell inspection limit.')
            raw_sheet = [list(row) for row in ws.iter_rows(values_only=True)]
            if norm(ws.title) not in RULE_SHEETS:
                result.update(_split_and_normalize_sheet_sections(ws.title, raw_sheet))
            else:
                result[ws.title] = raw_sheet
    finally:
        wb.close()
    return result


def candidate_headers(rows, bold_rows=None):
    bold_rows = bold_rows or set()
    candidates = []
    for number, row in enumerate(rows[:40], 1):
        filled = [v for v in row if v is not None and str(v).strip()]
        text = [v for v in filled if isinstance(v, str) and not v.startswith('=')]
        if len(text) < 2 or len(text) < len(filled) * 0.65:
            continue
        if len({norm(v) for v in text}) < 2:
            continue
        score = len(text) * 2 + (5 if number in bold_rows else 0) - number / 100
        if len({norm(v) for v in text}) != len(text):
            score -= 2
        # Penalize data rows containing emails, phone numbers, or specific Grade X values
        if any('@' in str(v) or re.fullmatch(r'(?:\+27|0)[1-9]\d{7,8}', re.sub(r'[\s()-]', '', str(v))) or re.fullmatch(r'grade\s*\d{1,2}[a-z]?', str(v).strip(), re.I) for v in text):
            score -= 18
        # Penalize banner/title/metadata rows so real column headers below them win (never penalize 'Date of Birth' or wide header rows with >=6 columns)
        if len(text) <= 5 and any(len(str(v).strip()) > 45 or ':' in str(v) or norm(v).startswith(('synthetic', 'fy20', 'report', 'prepared', 'confidential', 'summary —', 'summary -', 'page ', 'date:', 'school:', 'district:')) for v in text):
            score -= 8
        # Penalize total/subtotal summary rows so they are never picked as headers
        if norm(text[0]) in {'total', 'subtotal', 'grand total'}:
            score -= 10
        headers = [str(v).strip() if v is not None else '' for v in row]
        while headers and not headers[-1]:
            headers.pop()
        if all(h for h in headers):
            score += 2
        if number < len(rows):
            next_filled = [v for v in rows[number] if v is not None and str(v).strip()]
            if len(next_filled) >= max(2, len(headers) - 1):
                score += 3
        candidates.append({'row': number, 'headers': headers, 'score': score})
    return sorted(candidates, key=lambda item: item['score'], reverse=True)[:12]


def extract_rules(rows_by_sheet):
    config, aliases, rules, unknown = {}, {}, [], []
    for sheet, rows in rows_by_sheet.items():
        if norm(sheet) not in RULE_SHEETS:
            continue
        table_header = None
        for i, row in enumerate(rows[:15]):
            heads = [norm(v).replace('(s)', 's') for v in row[:4]]
            if len(heads) == 4 and heads[0] == 'id' and heads[1] in {'type', 'kind'} and heads[2] in {'fields', 'field'} and heads[3] == 'value':
                table_header = i
                break
        if table_header is not None:
            for number, row in enumerate(rows[table_header + 1:], table_header + 2):
                row = list(row) + [None] * max(0, 4 - len(row))
                if all(value is None for value in row[:4]):
                    continue
                rid, kind, field, value = (str(v).strip() if v is not None else '' for v in row[:4])
                kind = norm(kind)
                evidence = f'{sheet}!A{number}:D{number}'
                if kind == 'config' and field in CONFIG_KEYS:
                    if field in config and config[field] != value:
                        unknown.append({'ref': evidence, 'text': f'Conflicting config {field}: {config[field]} versus {value}'})
                    else:
                        config[field] = value
                elif kind == 'alias':
                    aliases[field] = list(dict.fromkeys([field] + [x.strip() for x in value.split('|') if x.strip()]))
                elif kind in KINDS:
                    fields = [x.strip() for x in field.split(',') if x.strip()]
                    if not fields:
                        unknown.append({'ref': evidence, 'text': f'Rule {rid} has no fields.'})
                    else:
                        rules.append(Rule(id=rid or evidence, kind=kind, fields=fields, value=value, evidence=evidence).model_dump())
                else:
                    unknown.append({'ref': evidence, 'text': ' | '.join([rid, kind, field, value])})
            continue
        # A narrow, documented declarative grammar; anything else remains an unresolved instruction.
        for number, row in enumerate(rows, 1):
            for column, cell in enumerate(row, 1):
                if not isinstance(cell, str) or not cell.strip():
                    continue
                text = cell.strip()
                ref = f'{sheet}!{get_column_letter(column)}{number}'
                if norm(text).startswith(('synthetic', 'rules —', 'rules -')) or norm(text) in {'rules', 'instructions'}:
                    continue
                match = re.match(r'^(mode|key|period|value|categories|total|header_row|output_sheet|write_mode)\s*:\s*(.+)$', text, re.I)
                if match:
                    config[match[1].lower()] = match[2].strip()
                    continue
                match = re.match(r'^(required|unique|nonnegative|numeric|integer)\s*:\s*(.+)$', text, re.I)
                if match:
                    rules.append(Rule(id=ref, kind=match[1].lower(), fields=[x.strip() for x in match[2].split(',')], evidence=ref).model_dump())
                    continue
                match = re.match(r'^sum\s*:\s*([^=]+)=\s*(.+)$', text, re.I)
                if match:
                    rules.append(Rule(id=ref, kind='sum', fields=[match[1].strip()], value=','.join(x.strip() for x in match[2].split('+')), evidence=ref).model_dump())
                    continue
                unknown.append({'ref': ref, 'text': text})
    # Duplicate IDs with different definitions must not overwrite one another.
    seen = {}
    for rule in rules:
        body = {k: v for k, v in rule.items() if k != 'evidence'}
        if rule['id'] in seen and seen[rule['id']] != body:
            unknown.append({'ref': rule['evidence'], 'text': f'Conflicting definitions for rule {rule["id"]}.'})
        seen[rule['id']] = body
    return config, aliases, rules, unknown


def profile(path: Path, file_id: str, template=False):
    rows_by_sheet = book_rows(path)
    has_rules_sheet = any(norm(s) in RULE_SHEETS for s in rows_by_sheet)
    config, aliases, rules, unknown = extract_rules(rows_by_sheet)
    bold, merges, formulas, unsupported = {}, {}, {}, []
    wb = None
    if path.suffix.lower() == '.xlsx':
        wb = load_workbook(path, data_only=False, keep_links=False)
        with zipfile.ZipFile(path) as archive:
            if any(x.startswith('xl/externalLinks/') for x in archive.namelist()):
                unsupported.append('External workbook links require explicit support; they are never fetched.')
        for ws in wb:
            bold[ws.title] = {cell.row for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)) for cell in row if cell.font.bold}
            merges[ws.title] = [str(r) for r in ws.merged_cells.ranges]
            formulas[ws.title] = [cell.coordinate for row in ws for cell in row if cell.data_type == 'f']
            if not template and has_rules_sheet and ws.data_validations.dataValidation:
                unsupported.append(f'Source native validation on {ws.title} needs input-stage compilation; this first version holds it rather than silently ignoring it.')
            if template and ws.protection.sheet and norm(ws.title) not in RULE_SHEETS:
                unsupported.append(f'Protected output sheet {ws.title} needs an explicit authorisation policy.')
    def capture_instruction(text, ref):
        cleaned = re.sub(r'^(?:rule|instruction|mapping rule|validation rule)\s*:\s*', '', text.strip(), flags=re.I)
        extra_config, extra_aliases, extra_rules, extra_unknown = extract_rules({'Instructions': [[cleaned]]})
        if extra_rules:
            for rule in extra_rules:
                rules.append(rule | {'id': ref, 'evidence': ref})
        elif extra_config:
            for key, value in extra_config.items():
                if key in config and config[key] != value:
                    unknown.append({'ref': ref, 'text': 'Conflicting workbook config: ' + cleaned})
                else:
                    config[key] = value
        else:
            unknown.append({'ref': ref, 'text': cleaned})

    for name, rows in rows_by_sheet.items():
        if norm(name) in RULE_SHEETS:
            continue
        guessed = candidate_headers(rows, bold.get(name))
        prefix_end = int(config.get('header_row') or (guessed[0]['row'] if guessed else 1)) - 1
        for number, row in enumerate(rows, 1):
            for col, value in enumerate(row, 1):
                if not isinstance(value, str) or norm(value).startswith('synthetic'):
                    continue
                explicit = re.match(r'^(rule|instruction|mapping rule|validation rule)\s*:', value, re.I)
                prefix_rule = has_rules_sheet and number <= prefix_end and re.search(r'\b(must|required|unique|rule|ignore|exclude|do not|nonnegative)\b', value, re.I)
                if explicit or prefix_rule:
                    capture_instruction(value, f'{name}!{get_column_letter(col)}{number}')
        if wb and has_rules_sheet and name in wb.sheetnames:
            for row in wb[name]:
                for cell in row:
                    if cell.comment and re.search(r'\b(must|required|unique|rule|ignore|exclude|do not|nonnegative)\b', cell.comment.text, re.I):
                        capture_instruction(cell.comment.text, f'{name}!{cell.coordinate} comment')
    unknown = list({(item['ref'], item['text']): item for item in unknown}.values())
    sheets = []
    for name, rows in rows_by_sheet.items():
        if norm(name) in RULE_SHEETS:
            continue
        candidates = candidate_headers(rows, bold.get(name))
        candidate = candidates[0] if candidates else {'row': 1, 'headers': [str(v or '') for v in (rows[0] if rows else [])], 'score': 0}
        if template and config.get('header_row'):
            number = int(config['header_row'])
            if number < 1 or number > len(rows):
                raise ValueError('Template header_row does not reference a real row.')
            candidate = {'row': number, 'headers': [str(v or '').strip() for v in rows[number - 1]], 'score': 100}
        headers = list(candidate['headers'])
        while headers and not headers[-1]:
            headers.pop()
        sheet = {
            'name': name, 'header_row': candidate['row'], 'headers': headers,
            'candidates': candidates, 'rows': len(rows), 'columns': max((len(r) for r in rows), default=0),
            'preview': [[public_value(v) for v in row] for row in rows[candidate['row']:candidate['row'] + 5]],
            'merged_ranges': merges.get(name, []), 'formula_cells': formulas.get(name, []),
        }
        sheet['layout_fingerprint'] = digest({'sheet': name, 'header_row': sheet['header_row'], 'headers': headers})
        sheets.append(sheet)
    if not sheets:
        if wb:
            wb.close()
        raise ValueError('No source or output data sheet was found.')
    selected = next((s for s in sheets if s['name'] == config.get('output_sheet')), max(sheets, key=lambda s: (len([h for h in s['headers'] if h]), s['rows'])) if not has_rules_sheet else sheets[0])
    if template and len(sheets) > 1 and not config.get('output_sheet') and has_rules_sheet:
        unsupported.append('Multiple possible output sheets: specify output_sheet in the Rules sheet.')
    if template and not has_rules_sheet and not config and not rules:
        clean_headers = [h for h in selected['headers'] if h] or ['Col1']
        period_cols = [h for h in clean_headers if norm(h) in PERIOD_TOKENS]
        total_col = next((h for h in clean_headers if norm(h) in {'total', 'sum', 'ytd', 'annual', 'grand total'}), '')
        if len(period_cols) >= 2:
            key_col = next((h for h in clean_headers if h not in period_cols and h != total_col), clean_headers[0])
            config.update({
                'mode': 'pivot', 'key': key_col, 'period': 'Period', 'value': 'Units',
                'categories': '|'.join(period_cols), 'total': total_col,
                'header_row': str(selected['header_row']), 'write_mode': 'replace',
                'ignore_row_labels': 'Subtotal|Total|Grand Total|Sum|Average|Notes|Source'
            })
            rules.append(Rule(id='AUTO-req', kind='required', fields=[key_col] + period_cols, evidence='Auto-detected table structure').model_dump())
            if total_col:
                rules.append(Rule(id='AUTO-sum', kind='sum', fields=[total_col], value=','.join(period_cols), evidence='Auto-detected Total column').model_dump())
            aliases['Period'] = builtin_aliases_for('Period')
            aliases['Units'] = builtin_aliases_for('Qty')
        else:
            key_col = clean_headers[0]
            config.update({
                'mode': 'direct', 'key': key_col,
                'header_row': str(selected['header_row']), 'write_mode': 'replace',
                'ignore_row_labels': 'Subtotal|Total|Grand Total|Sum|Average|Notes|Source'
            })
            rules.append(Rule(id='AUTO-req', kind='required', fields=[key_col], evidence='Auto-detected key column').model_dump())
        for h in clean_headers:
            aliases[h] = builtin_aliases_for(h)
    if template and wb:
        ws = wb[selected['name']]
        header_row = selected['header_row']
        for index, dv in enumerate(ws.data_validations.dataValidation):
            fields = set()
            for cell_range in dv.ranges.ranges:
                for col in range(cell_range.min_col, cell_range.max_col + 1):
                    if cell_range.max_row > header_row and col <= len(selected['headers']):
                        fields.add(selected['headers'][col - 1])
            fields = sorted(x for x in fields if x)
            if not fields:
                continue
            evidence = f'{ws.title}!{dv.sqref} native validation'
            if dv.type in {'whole', 'decimal'}:
                rules.append(Rule(id=f'N{index}-type', kind='integer' if dv.type == 'whole' else 'numeric', fields=fields, evidence=evidence).model_dump())
                ops = {'greaterThanOrEqual': [('min', dv.formula1)], 'lessThanOrEqual': [('max', dv.formula1)], 'between': [('min', dv.formula1), ('max', dv.formula2)]}
                if dv.operator in ops:
                    for kind, value in ops[dv.operator]:
                        try:
                            Decimal(str(value))
                            rules.append(Rule(id=f'N{index}-{kind}', kind=kind, fields=fields, value=str(value), evidence=evidence).model_dump())
                        except Exception:
                            unknown.append({'ref': evidence, 'text': 'Dynamic numeric validation bound is not yet supported.'})
                elif dv.operator:
                    unknown.append({'ref': evidence, 'text': f'Unsupported native validation operator {dv.operator}.'})
            elif dv.type == 'list' and dv.formula1 and dv.formula1.startswith('"'):
                values = dv.formula1.strip('"').split(',')
                rules.append(Rule(id=f'N{index}-allowed', kind='allowed', fields=fields, value='|'.join(values), evidence=evidence).model_dump())
            else:
                unknown.append({'ref': evidence, 'text': 'This native validation needs a supported compilation rule: ' + str(dv.type)})
    if wb:
        wb.close()
    family = digest({'sheet': selected['name'], 'headers': selected['headers'], 'mode': config.get('mode'), 'key': config.get('key')})
    version = digest({'family': family, 'config': config, 'aliases': aliases, 'rules': rules, 'unknown': unknown})
    return {
        'id': file_id, 'name': path.name, 'bytes': path.stat().st_size,
        'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': sheets,
        'config': config, 'aliases': aliases, 'rules': rules, 'unknown_rules': unknown,
        'unsupported_features': unsupported, 'output_sheet': selected['name'],
        'family': family, 'version': version,
    }
